# -*- coding: utf-8 -*-
"""덱을 고해상도 PDF 미리보기로 내보낸다 (150dpi, wqy 대체폰트)."""
import sys
import render_preview as R
from pptx import Presentation
from PIL import Image, ImageDraw

R.DPI = 150.0
R.SCALE = R.DPI / 914400.0
R.PT2PX = R.DPI / 72.0
R._fc.clear()

def export(pptx, out_pdf):
    prs = Presentation(pptx)
    W = R.emu_px(prs.slide_width); H = R.emu_px(prs.slide_height)
    pages = []
    for slide in prs.slides:
        bgc = (255, 255, 255)
        try:
            if slide.background.fill.type == 1:
                bgc = R.rgb(slide.background.fill.fore_color.rgb) or bgc
        except Exception: pass
        img = Image.new("RGB", (W, H), bgc)
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y = R.emu_px(sh.left or 0), R.emu_px(sh.top or 0)
            w, hh = R.emu_px(sh.width or 0), R.emu_px(sh.height or 0)
            fill = R.get_fill(sh); lc, lw = R.get_line(sh)
            try: st = sh.auto_shape_type
            except Exception: st = None
            if fill or lc:
                if st is not None and int(st) == 9:
                    d.ellipse([x, y, x + w, y + hh], fill=fill, outline=lc, width=lw or 1)
                elif st is not None and int(st) == 5:
                    r = max(2, int(min(w, hh) * 0.12))
                    d.rounded_rectangle([x, y, x + w, y + hh], radius=r, fill=fill, outline=lc, width=lw or 1)
                else:
                    d.rectangle([x, y, x + w, y + hh], fill=fill, outline=lc, width=lw or 1)
            if sh.has_text_frame and sh.text_frame.text.strip():
                R.draw_text_frame(d, sh, x, y, w, hh)
        pages.append(img)
    pages[0].save(out_pdf, "PDF", save_all=True, append_images=pages[1:], resolution=150.0)
    print("PDF saved:", out_pdf, "pages:", len(pages))

if __name__ == "__main__":
    export(sys.argv[1], sys.argv[2])
