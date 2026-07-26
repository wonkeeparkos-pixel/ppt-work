# -*- coding: utf-8 -*-
"""보툴리눔 덱 — 정적 PPTX(16:9 풀블리드) + 미리보기 PDF.

각 웹 슬라이드를 PNG로 래스터화한 결과를 그대로 배치한다.
임베드 미디어/타이밍이 없어 어떤 PowerPoint에서도 경고 없이 열리고 저장된다.
"""
import os
import glob
import sys
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

BASE = "/home/user/ppt-work/문헌고찰_보툴리눔_정형외과통증"
PNG_DIR = sys.argv[1] if len(sys.argv) > 1 else (
    "/tmp/claude-0/-home-user-ppt-work/4a16da38-a82b-5d5e-a9c9-97ae114e4ff2/scratchpad/png_botox"
)
PPTX = os.path.join(BASE, "보툴리눔_정형외과통증_발표.pptx")
PDF = os.path.join(BASE, "보툴리눔_정형외과통증_발표_미리보기.pdf")

pngs = sorted(glob.glob(os.path.join(PNG_DIR, "s*.png")))
assert pngs, f"no png in {PNG_DIR}"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for png in pngs:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save(PPTX)
print(f"pptx: {PPTX} | {os.path.getsize(PPTX)/1e6:.1f} MB | {len(pngs)} slides")

pages = [Image.open(p).convert("RGB") for p in pngs]
pages[0].save(PDF, save_all=True, append_images=pages[1:], resolution=150.0)
print(f"pdf : {PDF} | {os.path.getsize(PDF)/1e6:.1f} MB")
