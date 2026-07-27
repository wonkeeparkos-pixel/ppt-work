# -*- coding: utf-8 -*-
"""웹 발표본(LSB_발표_웹.html)을 '웹 디자인 그대로' PPTX로 변환.
각 웹 슬라이드(.stage, 16:9)를 고해상도 PNG로 렌더한 뒤, 13.333x7.5in 슬라이드에 풀블리드로 배치.

렌더 단계(Chromium, deviceScaleFactor=3)는 아래 node 스크립트로 생성한다(예: /tmp/hires.js):
  각 .snap .stage 를 순회하며 hs-01.png ... hs-12.png 저장(3600x2028px).
이 스크립트는 그 PNG들을 모아 PPTX로 조립한다(이미지 기반 = 픽셀 정확, 텍스트 편집은 불가).
텍스트 편집이 필요하면 네이티브 빌더 build_lsb.py(=LSB_발표.pptx)를 사용."""
import os, sys, glob
from pptx import Presentation
from pptx.util import Inches

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_DIR = sys.argv[1] if len(sys.argv) > 1 else os.path.join(REPO, "build", "hires")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
    REPO, "문헌고찰_NLC_RLS_LSB", "03_LSB", "LSB_발표_웹.pptx")

imgs = sorted(glob.glob(os.path.join(IMG_DIR, "hs-*.png")))
if not imgs:
    sys.exit("no rendered slides in %s — run the Chromium render first" % IMG_DIR)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
for p in imgs:
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.save(OUT)
print("web->pptx:", len(imgs), "slides ->", OUT, f"({os.path.getsize(OUT)/1024/1024:.1f}MB)")
