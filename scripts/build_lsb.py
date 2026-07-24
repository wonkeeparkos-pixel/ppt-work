# -*- coding: utf-8 -*-
"""LSB(요추교감신경차단) 발표 덱 전용 빌더 — LSB 폴더 자료 기반."""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from ppt_lib import Deck, INK, TEALD, GREEN, MINT, RED, MUTE, AMBER, WHITE

_ASSET = "/home/user/ppt-work/문헌고찰_NLC_RLS_LSB/03_LSB/assets/"
AXIAL_IMG = _ASSET + "lsb_axial.png"
CONTRAST_IMG = _ASSET + "lsb_contrast.png"
GF_IMG = _ASSET + "lsb_gf.png"

def fig_slide(d, title, eyebrow, tag, img, caption, foot, h=4.5):
    s = d._slide()
    d.rect(s, 0, 0, d.SW, d.SH, WHITE)
    d.header(s, title, eyebrow, tag)
    pic = s.shapes.add_picture(img, 0, Inches(1.55), height=Inches(h))
    pic.left = int((d.SW - pic.width) / 2)
    d.text(s, Inches(0.7), Inches(6.4), Inches(11.95), Inches(0.5),
           [[(caption, 11, TEALD, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    d.footer(s, foot)
    return s

BASE = "/home/user/ppt-work/문헌고찰_NLC_RLS_LSB"

d = Deck()
d.title_slide("요추교감신경차단 (Lumbar Sympathetic Block)",
    ["방법 · 적응증 · 효과", ""],
    "Lumbar Sympathetic Block / Sympatholysis — Technique, Indications, and Efficacy",
    "구성: 개요/해부 → 방법 → 적응증 → 효과 → 합병증 → 결론",
    "교육·연구 참고용. 개별 환자의 진료 결정은 담당 의사의 판단에 따른다.  |  작성 2026-07")

# ---------- 감별 오프너 (NLC 스타일: 감별이 먼저) ----------
d.table_slide("밤·하지 증상 감별 — LSB는 어디에 쓰나", "가장 먼저", "감별", [
    ("질환","핵심 소견","LSB 적응"),
    ("허혈성 안정통 (PAD·CLI)","파행 → 야간 안정통·차고 창백한 발·ABI↓ (동맥성)","○"),
    ("하지 CRPS","작열통·이질통·부종·피부색/온도 좌우차 (교감매개통)","○"),
    ("당뇨병성 신경병증","저림·화끈거림·감각이상·야간 악화","△ 난치성"),
    ("특발성 NLC (야간경련)","통증성 근수축·촉지 근경직·족배굴곡 완화","✕"),
    ("RLS (하지불안)","움직임 충동·움직이면 완화·근경직 없음","✕"),
    ("정맥부전·정맥류","부종·무거움·저녁 악화 (정맥 역류)","✕"),
], [Inches(2.9),Inches(7.15),Inches(1.8)], "감별이 치료의 출발점 · ABI·도플러 선행",
    tag_color=INK, size=12, row_h=Inches(0.6),
    note="핵심: ABI·도플러로 동맥성 여부 감별이 출발점 — 정맥류·특발경련·RLS는 LSB 적응 아님.")

d.bullets_slide("개요 · 해부 · 원리", "개요", "배경", [
    (0,"LSB와 교감신경 신경파괴(sympatholysis)는 70년 이상 다양한 하지 통증·허혈 질환에 사용.",INK,True),
    (0,"표적: 신경절 밀도가 가장 높은 L2–L3(보통 'L2 하 1/3 ~ L3 상 1/3'), 척추체 전외측. 일부 L2/L3/L4 다분절.",INK),
    (0,"원리 ①: 교감신경 차단 → 혈관 확장·측부순환 증가 → 조직 산소화 개선.",INK),
    (0,"원리 ②: 교감신경 매개 통증(sympathetically-maintained pain) 경로 및 자율신경 동반 침해 구심로 차단.",INK),
    (-1,"국소마취제 차단(진단적/치료적)과 화학적/열적 신경파괴로 나뉜다.",TEALD,True),
], "StatPearls NBK431107; Zhang 2022(Ibrain)")

d.split_slide("방법 (Technique)", "방법", "기법", [
    (0,"표준: 방척추(paravertebral) 접근 + 투시(fluoroscopy). CT·초음파도 가능.",INK,True),
    (0,"바늘 진입: 정중선에서 약 7 cm 외측. 척추체 접촉 후 전내측으로 'walk'하여 척추체 전외측으로 진입.",INK),
    (0,"흡인 후 조영제 주입 → 두미측(craniocaudal) 종방향 확산 확인.",INK),
    (0,"성공 지표: 동측 하지 피부온도 ≥2°C 상승.",INK,True),
], ("약제", [
    (0,"진단/치료: lidocaine 1%,",None),
    (0,"  bupivacaine 0.25–0.5%, ropivacaine",None),
    (0,"신경파괴: 무수알코올/phenol, RFA",None),
    (0,"원칙: 신경파괴는 진단적 차단 양성 시에만",None),
], None), "StatPearls; Lumbar Sympatholysis NBK560514", tag_color=TEALD)

# ---------- L2·L3 조감도 (오리지널 도해 이미지) ----------
_s = d._slide()
d.rect(_s, 0, 0, d.SW, d.SH, WHITE)
d.header(_s, "L2·L3 조감도 — 표적과 위험 구조", "방법 · 그림으로", "축상면 axial")
_pic = _s.shapes.add_picture(AXIAL_IMG, 0, Inches(1.45), height=Inches(4.55))
_pic.left = int((d.SW - _pic.width) / 2)
d.text(_s, Inches(0.7), Inches(6.3), Inches(11.95), Inches(0.5),
       [[("표적=척추체 전외측 교감신경절 · 방척추(정중선 ~7cm) 접근 · 대동맥·IVC·요관·신장·생식대퇴신경·추간공 회피", 11, TEALD, True)]],
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
d.footer(_s, "교육용 오리지널 도해 · StatPearls NBK431107")

# ---------- L2·L3 시술 주의점 ----------
d.bullets_slide("L2·L3 시술 — 전·중 주의할 점", "방법 · 안전", "시술 주의", [
    (0,"영상 유도 필수: 투시(또는 CT). 조영제로 두미측 종방향 확산 확인 — 후방(추간공)·혈관 확산 시 즉시 재위치.",INK),
    (0,"바늘 끝은 척추체 전외측에 — psoas·추간공 진입 금지.",INK,True),
    (0,"레벨은 L2 하1/3~L3 상1/3, L4로 내려가지 말 것(생식대퇴신경통 급증).",INK,True),
    (0,"흡인 후 분할·점진 주입(혈관내·경막외 조기 발견) · 소량 시험주입.",INK),
    (0,"항응고·항혈소판제 중단(심부·후복막 차단 — 출혈 위험) · 무균술.",INK),
    (-1,"성공지표 온도 ≥2°C↑ 확인 · 시술 후 혈압·하지 근력/감각 모니터.",TEALD,True),
], "StatPearls NBK431107 · NBK557637 · Feigl 1998(PMID 9425975)", size=14, gap=7)

# ---------- 조영제 확산 읽기 (도해 이미지) ----------
fig_slide(d, "조영제 확산 읽기 — 이 패턴이면 멈춰라", "안전 · 핵심기술", "조영제 읽기", CONTRAST_IMG,
    "실시간 조영제가 최고의 안전장치 — 종방향(위·아래) 확산만 진행 · 혈관·경막외·근육내 패턴이면 즉시 재위치",
    "조영제 확산 패턴 판독 · StatPearls NBK431107", h=4.35)

# ---------- 합병증 & 회피법 (표) ----------
d.table_slide("합병증 & 회피법", "안전성", "합병증·회피", [
    ("합병증","원인·기전","피하는 법"),
    ("생식대퇴신경통 (5–10%)","psoas 역류·자극 (L2 0%·L4 40%)","L4 회피·psoas 주입금지·최소용량"),
    ("혈관손상·혈관내주입","대동맥/IVC·요추혈관 인접","투시+흡인+조영제·항응고 중단"),
    ("요관·신장 천공","후복막 장기 인접","조영제로 깊이·궤적 확인"),
    ("체성신경·신경축 확산","바늘 후방→경막외 tracking","바늘 전외측 유지"),
    ("기립성 저혈압","교감차단 → 혈관확장","수액·서서히 기립·양측 신중"),
    ("신경염·사정장애","알코올/phenol·양측 L1–L2","진단차단 양성 시만·L1–L2 회피"),
], [Inches(3.0),Inches(4.55),Inches(4.3)], "StatPearls NBK431107·NBK557637 · Feigl 1998(PMID 9425975)",
    tag_color=RED, size=11.5, row_h=Inches(0.62),
    note="핵심: 투시+조영제로 확산 확인 · 바늘 전외측 유지 · L2~L3 표적 · 신경파괴는 진단차단 양성 시.")

# ---------- 생식대퇴신경통 심화 (도해 이미지) ----------
fig_slide(d, "생식대퇴신경통 — 가장 흔한 걱정, 대부분 회피 가능", "합병증 · 심화", "최다 우려", GF_IMG,
    "회피: L2하~L3상 표적·L4 회피·psoas 주입금지·소량 / 발생 시: 대개 수 주 내 자연호전·신경병증통 약물로 대증",
    "Feigl 1998(PMID 9425975); StatPearls NBK431107", h=4.5)

# ---------- 혈관·LAST·신경 손상 대응 ----------
d.bullets_slide("혈관·LAST·신경 손상 — 조기 인지와 대응", "합병증 · 대응", "인지·구조", [
    (0,"LAST(국소마취제 전신독성): 입 주변 저림·이명·금속맛·어지럼 → 경련·부정맥.",RED,True),
    (0,"대응 — 즉시 중단·산소/기도·경련조절·20% 지질유탁액 정주·소생술. 예방 — 흡인+조영제+분할주입+용량제한.",INK),
    (0,"신경·신경축: 주입 중 방사통·하지 위약 → 즉시 중단·재위치. 시술 후 근력·감각 확인.",INK),
    (0,"출혈·감염: 항응고 중단·무균술 · 심한 통증·발열·팽창 시 즉시 평가.",INK),
    (-1,"준비된 상태(모니터·정맥로·지질유탁액)면 대부분 안전하게 관리된다.",TEALD,True),
], "ASRA LAST 지침 · StatPearls NBK431107", tag_color=RED, size=14, gap=8)

# ---------- 합병증 대비 · 환자설명 ----------
d.bullets_slide("합병증 대비 — 준비·조기인지·환자설명", "합병증 · 대비", "환자 안심", [
    (0,"대부분 경미·자가회복 — 심각 합병증은 영상유도·정확한 술기로 드묾.",INK,True),
    (0,"준비: 모니터·정맥로·응급장비·20% 지질유탁액 구비, 소생 프로토콜 숙지.",INK),
    (0,"조기 인지: 조영제 패턴 + 환자 증상(통증·저림·어지럼)을 실시간 관찰.",INK),
    (0,"환자 설명(동의): 흔한 것(일시적 저림·주사부위통·저혈압) vs 드문 것(신경통·출혈)을 사전 고지 → 신뢰·불안↓.",INK),
    (-1,"신경파괴는 진단차단 양성 시에만·소량 — 위해를 최소화.",TEALD,True),
], "교육용 정리", tag_color=INK, size=14, gap=8)

d.table_slide("적응증 (Indications)", "적응증", "적응", [
    ("범주","대표 적응증"),
    ("통증증후군","하지 복합부위통증증후군(CRPS I/II) — 조기(≤12개월) 시행이 유리"),
    ("허혈질환","PAD·중증 하지허혈(CLI) 안정통·비재건성, 버거병, 색전, 동상, 혈관연축"),
    ("신경병증","당뇨병성 신경병증/당뇨발, 대상포진후신경통, 환상지통·단단통"),
    ("기타","족부 다한증, 레이노 증후군(하지)·홍색사지통증, 암성 골반/하지통"),
], [Inches(2.4),Inches(9.45)], "StatPearls; 문헌고찰 종합", size=12, row_h=Inches(0.7),
    note="특발성 야간 하지경련(NLC)은 적응증 아님 — '밤 다리증상'이 허혈성 안정통(PAD/CLI)이면 적응. 원인 감별(ABI·도플러) 선행.")

# ---------- 적응증별 환자 증상·호소 ① ----------
d.bullets_slide("적응증별 환자 증상·호소 ① — 통증증후군·허혈", "적응증 · 증상", "환자 증상", [
    (0,"하지 CRPS I/II (복합부위통증증후군) — 교감신경 매개 통증의 전형",INK,True),
    (1,"작열통·이질통(옷·바람 스침에도 통증)·통각과민 · 부종 · 피부색(발적↔창백)·좌우 온도차 · 발한 이상",INK),
    (1,"호소: “발이 타는 듯 화끈거려요” · “이불만 스쳐도 아파요” · “붓고 색이 자꾸 변해요”",MUTE),
    (0,"허혈질환 — PAD·중증 하지허혈(CLI)·버거병 (동맥성 허혈)",GREEN,True),
    (1,"간헐적 파행(걷다 종아리 통증→쉬면 완화) → 진행 시 야간 안정통(다리 내리면 완화) · 차고 창백한 발 · 비치유 궤양·괴저",INK),
    (1,"호소: “걸으면 종아리가 터질 듯 아파요” · “밤에 발이 시려 잠을 못 자요” · “상처가 안 아물어요”",MUTE),
    (-1,"‘허혈’은 동맥성 부족을 의미 — 하지정맥류 등 정맥질환은 LSB 적응증 아님(압박·정맥폐색술이 표준).",TEALD,True),
], "StatPearls; 문헌고찰 종합", size=14, gap=6)

# ---------- 적응증별 환자 증상·호소 ② ----------
d.bullets_slide("적응증별 환자 증상·호소 ② — 신경병증·기타", "적응증 · 증상", "환자 증상", [
    (0,"신경병증 — 당뇨병성 신경병증/당뇨발·대상포진후신경통·환상지통",INK,True),
    (1,"저림·화끈거림·전기 오듯 찌름 · 이질통 · 감각저하(양말 신은 느낌) · 야간 악화 · 당뇨발 궤양·관류저하",INK),
    (1,"호소: “발이 저리고 화끈거려요” · “밤에 더 심해요” · “전기가 찌릿 오는 것 같아요”",MUTE),
    (0,"기타 — 족부 다한증·레이노(하지)·홍색사지통증·암성 하지통",INK,True),
    (1,"다한증(과한 발땀·축축·악취) · 레이노(추위·스트레스에 창백→청색→발적 삼색변화·저림) · 홍색사지통증(발작적 발적·작열·열감, 열에 악화)",INK),
    (1,"호소: “발에 땀이 너무 많아요” · “추우면 발이 하얘졌다 파래져요” · “발이 화끈 달아올라요”",MUTE),
], "StatPearls; 문헌고찰 종합", size=14, gap=6)

# ---------- 야간 신경병증통 서사 + LSB 치료 삽입 (NLC 스타일) ----------
d.split_slide("밤에 저리고 화끈거리는 다리 — 당뇨병성 신경병증과 LSB", "적응증 · 심화", "야간 신경병증통", [
    (0,"임상상: 당뇨병성 말초신경병증 — 발·종아리 저림·화끈거림·전기 찌름, 양측 ‘양말’ 분포.",INK,True),
    (0,"야간통: 이불 온기·야간 순환·주의분산 소실 → 밤에 증폭·수면 방해.",INK),
    (0,"먼저 감별: RLS(움직이면 완화)·NLC(경련)·허혈성 안정통과 구분.",INK),
    (-1,"1차: 혈당조절+약물(가바펜틴·듀록세틴) → 난치성이면 LSB 고려",TEALD,True),
], ("LSB — 치료 삽입", [
    (0,"기전: 교감차단 → 미세순환↑ + 교감매개통 차단",None),
    (0,"근거: 난치성 DPN에 LSB+신경파괴 RCT·증례",None),
    (0,"적용: 약물 불응·진단차단 양성 시 선택적 시행",None),
    (0,"성공지표: 피부온도 ≥2°C↑ (근거 제한적)",None),
], None), "Zhang 2020(RCT) PMID 32915421; 증례 PMID 22606406; StatPearls NBK442009", dark_card=True)

d.split_slide("효과 (1) 관류 개선 · 허혈성 통증", "효과", "Lv III–IV", [
    (0,"관류 실측(Dickey 2024): LSB 후 후경골동맥 직경 0.17→0.27 cm(+58.8%), 모세혈관 재충혈 3.92→1.30초, 족부온도 +2.8°C.",INK,True),
    (0,"허혈성 통증(PAD): 하지 통증 최대 75%↓, Fontaine 분류·측부관류 개선(Medicina 2024).",GREEN),
    (0,"비재건성 CLI: 교감신경 신경파괴가 절단 외 대안이 없는 환자의 통증조절에 효과적·안전(Barreto 2018).",INK),
], ("실측 수치(Dickey 2024)", [
    (0,"후경골동맥 +58.8%",MINT,True),
    (0,"모세혈관 재충혈 3.92→1.30초",None),
    (0,"족부온도 +2.8°C",None),
    (0,"→ 관류 개선의 객관 근거",MINT,True),
], None), "Dickey 2024(Cureus); Medicina 2024; Barreto 2018", tag_color=AMBER, dark_card=True)

d.bullets_slide("효과 (2) CRPS · 기전", "효과", "Lv III–IV", [
    (0,"CRPS는 교감신경 매개 통증의 전형 → LSB 후 유의한 통증완화가 축적된 근거. 최적 환자 선택·조기 시행이 성공률↑.",INK,True),
    (0,"Choi 2024(Sci Rep): 교감신경 신경파괴의 효과 지속기간을 전향 관찰로 제시.",INK),
    (0,"반응 예측: 교감신경 피부반응(sympathetic skin response)이 LSB 반응 예측에 유용(Pain Ther 2023).",INK),
    (0,"한계: 중추 감작이 진행되면 시간이 지날수록 효과 감소 가능.",MUTE),
    (-1,"기전 요약: 교감차단 → 측부순환 혈관확장 → 조직 산소화↑ → 통증↓ + 교감매개통 차단 + 신경파괴 직접효과.",TEALD,True),
], "Choi 2024(Sci Rep); Pain Ther 2023", tag_color=AMBER)

d.bullets_slide("근거수준 · 결론", "안전성 · 결론", "결론", [
    (0,"합병증·회피는 앞의 조감도·회피표 참고 — 대부분 영상유도·정확한 바늘 위치로 예방 가능.",RED,True),
    (0,"근거수준: 대부분 관찰·증례군·소규모 전향연구로 대규모 RCT는 부족(허혈·CRPS Level III~IV). 단, 난치성 당뇨병성 신경병증엔 RCT 존재.",INK),
    (0,"시행 원칙: 원인 감별 → 진단적 국소마취제 차단(온도 ≥2°C 확인) → 반응 양호 시 신경파괴/RFA.",INK),
    (-1,"그럼에도 조기 CRPS·재건 불가능한 중증 하지허혈(안정통)·난치성 신경병증에서 임상적으로 유용.",TEALD,True),
], "StatPearls; 문헌고찰 종합", tag_color=RED)

# ---------- 나의 프로토콜 · 실제 적용 (NLC 스타일 단계적 접근) ----------
d.bullets_slide("실제 진료 순서 — 교감신경 축을 겨냥한 단계적 접근", "나의 프로토콜 · 실제 적용", "실전 순서", [
    (0,"① 감별·검사: '밤/하지 증상'이 동맥성 허혈(ABI·도플러)·CRPS·신경병증인지 확인. 특발 경련·RLS·정맥류 배제.",INK),
    (0,"② 방향 설정: 교감매개통·허혈로 판단되면 교감신경 축(L2–L3)을 겨냥.",INK),
    (0,"③ 진단적 차단: 투시 유도 국소마취제 LSB → 피부온도 ≥2°C 상승·통증 반응 확인.",GREEN,True),
    (0,"④ 반응 양호 시: 치료적 반복 차단, 또는 신경파괴(무수알코올/phenol)·RFA로 장기 완화.",GREEN,True),
    (0,"⑤ 대상별: CRPS는 조기(≤12개월) 유리 · 난치성 당뇨병성 신경병증은 지속 LSB+신경파괴(RCT).",INK),
    (-1,"⑥ 주의: 특발성 NLC·정맥질환은 적응 아님 · 합병증(생식대퇴신경통 5–10% 등) 사전 고지.",TEALD,True),
], "교육용 제안 · 개별 적용은 임상 판단 · 신경파괴는 진단차단 양성 시에만", tag_color=INK, size=14, gap=8)

d.key_slide("핵심 메시지 — LSB", "표준은 투시 유도 방척추 접근, 성공은 온도 ≥2°C", [
    ("방법","L2–L3 방척추 접근·투시, 정중선 7 cm 외측, 조영제 두미측 확산, 성공지표 온도 ≥2°C."),
    ("약제","진단/치료는 국소마취제, 신경파괴는 무수알코올/phenol·RFA — 진단차단 양성 시에만."),
    ("적응증","조기 CRPS·PAD/CLI 안정통·신경병증·다한증 등. 특발성 NLC·정맥질환(정맥류)은 적응 아님."),
    ("효과","관류 실측 개선(Dickey), 허혈통 최대 75%↓, CRPS 통증완화(교감피부반응으로 예측)."),
    ("근거","대규모 RCT 부족(관찰·증례 중심). 조기 CRPS·비재건성 CLI에서 유용."),
])

d.refs_slide("참고문헌 — LSB", [
 "Dua A, Varacallo MA. Lumbar sympathetic block. StatPearls. 2026. NBK431107.",
 "Lumbar sympatholysis. StatPearls. NBK560514.",
 "Zhang JH, Deng YP, Geng MJ. Efficacy of the lumbar sympathetic ganglion block in lower limb pain. Ibrain. 2022;8(4):442-52. PMID 37786587.",
 "Dickey Z, Sharma N. Lumbar sympathetic block leading to increased arterial diameter and blood flow. Cureus. 2024;16(6):e61755. PMID 38975506.",
 "Barreto Junior EPS, et al. Neurolytic block of the lumbar sympathetic chain in critical lower limb ischemia. Braz J Anesthesiol. 2018;68(1):100-3. PMC9391669.",
 "Choi EJ, et al. Effect duration of lumbar sympathetic ganglion neurolysis in CRPS. Sci Rep. 2024;14(1):12693. PMID 38830944.",
 "Gungor S, et al. Sympathetic blocks for CRPS: a case series. Medicine (Baltimore). 2018;97(19):e0705. PMID 29742728.",
 "Effect of lumbar sympathetic blockade on pain, Fontaine classification, collateral perfusion in PAD. Medicina (Kaunas). 2024;60(5):682.",
 "Prediction of LSB efficacy in CRPS type 1 by sympathetic skin response. Pain Ther. 2023;12(3):809-22. PMC10199976.",
 "Continuous LSB + sympatholysis for refractory painful diabetic neuropathy — RCT. 2020. PMID 32915421.",
 "Sympathetic blocks: sustained relief in refractory painful diabetic neuropathy (case). 2012. PMID 22606406.",
 "Lumbar sympathectomy for ischaemia·vasculitis·diabetic neuropathy·hyperhidrosis — series. 2018. PMID 29516399.",
])

out = os.path.join(BASE, "03_LSB", "LSB_발표.pptx")
n = d.save(out)
print("LSB", n, "slides ->", out)
