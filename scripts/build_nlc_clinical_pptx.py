# -*- coding: utf-8 -*-
"""임상 원장(Clinical Ledger) 디자인의 NLC 발표 — 편집용 단일 .pptx (Pretendard)."""
import sys, os
sys.path.insert(0, "/home/user/ppt-work/scripts")
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --- load pre-modification NLC content ---
ns = {'__file__': '/tmp/bwd_data.py'}
exec(open('/tmp/bwd_data.py').read(), ns)
NLC = ns['NLC']

FONT = "Pretendard"
PAPER=RGBColor(0xF7,0xF8,0xF6); INK=RGBColor(0x17,0x23,0x2C); TEAL=RGBColor(0x0E,0x7C,0x7B)
TEALD=RGBColor(0x0B,0x5F,0x5E); MUTE=RGBColor(0x6B,0x76,0x80); BODY=RGBColor(0x33,0x40,0x3B)
LINE=RGBColor(0xE4,0xE9,0xE8); CARDBG=RGBColor(0xF0,0xF4,0xF3); CARDLN=RGBColor(0xDD,0xE6,0xE4)
GREEN=RGBColor(0x2E,0x7D,0x32); AMBER=RGBColor(0x9A,0x68,0x00); RED=RGBColor(0xA8,0x35,0x2A)
WHITE=RGBColor(0xFF,0xFF,0xFF); MINT=RGBColor(0x5F,0xD6,0xCC); ICE=RGBColor(0xBF,0xD3,0xE6)
LGRAY=RGBColor(0xEF,0xF3,0xF2)
TAGCOL={'':TEAL,'green':GREEN,'amber':AMBER,'red':RED,'ink':INK}
CLSCOL={'':BODY,'accent':TEALD,'green':GREEN,'red':RED,'muted':MUTE}

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]; SW,SH=prs.slide_width,prs.slide_height

def slide(): return prs.slides.add_slide(BLANK)
def setf(r,size,color,bold=False,italic=False,font=FONT):
    r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; r.font.italic=italic; r.font.name=font
    rPr=r._r.get_or_add_rPr()
    ea=rPr.find(qn('a:ea'))
    if ea is None: ea=rPr.makeelement(qn('a:ea'),{}); rPr.append(ea)
    ea.set('typeface',font)
def rect(s,x,y,w,h,fill,line=None,lw=1.0,round_=False):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,x,y,w,h)
    if round_:
        try: shp.adjustments[0]=0.09
        except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    shp.shadow.inherit=False; return shp

def _segs(t):
    out=[];bold=False;buf='';i=0
    while i<len(t):
        if t.startswith('<b>',i):
            if buf:out.append((buf,bold));buf='';
            bold=True;i+=3
        elif t.startswith('</b>',i):
            if buf:out.append((buf,bold));buf='';
            bold=False;i+=4
        else: buf+=t[i];i+=1
    if buf:out.append((buf,bold))
    return out

def para(p,text,size,base,boldcolor=None,link=None):
    boldcolor=boldcolor or INK
    for seg,b in _segs(text):
        r=p.add_run(); r.text=seg; setf(r,size,(boldcolor if b else base),b)
        if link:
            r.hyperlink.address=link
def tb(s,x,y,w,h,anchor=MSO_ANCHOR.TOP,wrap=True):
    t=s.shapes.add_textbox(x,y,w,h);tf=t.text_frame;tf.word_wrap=wrap;tf.vertical_anchor=anchor
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;return tf

def header(s,eyebrow,title,tag=None):
    if eyebrow:
        tf=tb(s,Inches(0.8),Inches(0.42),Inches(9),Inches(0.3))
        p=tf.paragraphs[0];r=p.add_run();r.text=eyebrow.upper();setf(r,11,TEAL,True)
    tf=tb(s,Inches(0.8),Inches(0.74),Inches(10.3),Inches(0.7),anchor=MSO_ANCHOR.MIDDLE)
    p=tf.paragraphs[0];r=p.add_run();r.text=title;setf(r,24,INK,True)
    rect(s,Inches(0.8),Inches(1.42),Inches(11.73),Pt(1.4),LINE)
    if tag:
        lbl,cc=tag; col=TAGCOL.get(cc,TEAL); w=Inches(0.34+0.115*len(lbl))
        rect(s,SW-w-Inches(0.8),Inches(0.5),w,Inches(0.4),col,round_=True)
        tf=tb(s,SW-w-Inches(0.8),Inches(0.5),w,Inches(0.4),anchor=MSO_ANCHOR.MIDDLE)
        p=tf.paragraphs[0];p.alignment=PP_ALIGN.CENTER;r=p.add_run();r.text=lbl;setf(r,11.5,WHITE,True)

def footer(s,src,pg,dark=False):
    c=MUTE if not dark else RGBColor(0x9F,0xB0,0xAD)
    rect(s,Inches(0.8),Inches(6.86),Inches(11.73),Pt(1.2),LINE if not dark else RGBColor(0x2A,0x39,0x42))
    tf=tb(s,Inches(0.8),Inches(6.96),Inches(10.6),Inches(0.34),anchor=MSO_ANCHOR.MIDDLE)
    p=tf.paragraphs[0];r=p.add_run();r.text="근거: "+src;setf(r,9.5,c)
    tf=tb(s,Inches(12.0),Inches(6.96),Inches(0.9),Inches(0.34),anchor=MSO_ANCHOR.MIDDLE)
    p=tf.paragraphs[0];p.alignment=PP_ALIGN.RIGHT;r=p.add_run();r.text=f"{pg:02d}";setf(r,11,c,True)

def bullets(s,x,y,w,h,items,size=15,gap=8):
    tf=tb(s,x,y,w,h)
    for i,it in enumerate(items):
        lvl=it[0];tx=it[1];cls=it[2] if len(it)>2 else ''
        base=CLSCOL.get(cls,BODY); bc=base if cls else INK
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.space_after=Pt(gap);p.line_spacing=1.12;p.alignment=PP_ALIGN.LEFT
        if lvl==0:
            r=p.add_run();r.text="●  ";setf(r,9,TEAL,True)
        elif lvl>=1:
            r=p.add_run();r.text="      –  ";setf(r,size,MUTE,True)
        para(p,tx,size,base,bc)

def stat_row(s,x,y,w,items):
    n=len(items); cw=w/n
    for i,(big,lbl) in enumerate(items):
        tf=tb(s,x+Emu(int(cw*i)),y,Emu(int(cw))-Inches(0.1),Inches(1.0))
        p=tf.paragraphs[0];r=p.add_run();r.text=big;setf(r,26,TEAL,True)
        p2=tf.add_paragraph();p2.space_before=Pt(2);r=p2.add_run();r.text=lbl;setf(r,11,MUTE,True)

def render(s_dict,pg):
    T=s_dict['t']; s=slide(); rect(s,0,0,SW,SH,PAPER)
    if T=='title':
        rect(s,0,0,Inches(0.2),SH,TEAL)
        tf=tb(s,Inches(0.95),Inches(1.55),Inches(11),Inches(0.35))
        r=tf.paragraphs[0].add_run();r.text=s_dict['eyebrow'].upper();setf(r,13,TEAL,True)
        tf=tb(s,Inches(0.95),Inches(2.0),Inches(11.4),Inches(1.4))
        r=tf.paragraphs[0].add_run();r.text=s_dict['title'];setf(r,44,INK,True)
        rect(s,Inches(0.97),Inches(3.5),Inches(4.2),Pt(1.6),LINE)
        tf=tb(s,Inches(0.95),Inches(3.75),Inches(11),Inches(0.5))
        r=tf.paragraphs[0].add_run();r.text=s_dict['sub'];setf(r,20,RGBColor(0x3A,0x47,0x50))
        tf=tb(s,Inches(0.95),Inches(4.4),Inches(11),Inches(0.4))
        r=tf.paragraphs[0].add_run();r.text=s_dict['order'];setf(r,14,TEAL,True)
        rect(s,Inches(0.95),Inches(6.5),Inches(11.4),Pt(1.2),LINE)
        tf=tb(s,Inches(0.95),Inches(6.62),Inches(11),Inches(0.34),anchor=MSO_ANCHOR.MIDDLE)
        r=tf.paragraphs[0].add_run();r.text=s_dict['series'];setf(r,10.5,MUTE)
        return
    header(s,s_dict.get('eyebrow'),s_dict['title'],s_dict.get('tag'))
    if T=='bullets':
        has_stat=bool(s_dict.get('stat')); has_note=bool(s_dict.get('note'))
        bh=Inches(4.0) if (has_stat or has_note) else Inches(5.0)
        bullets(s,Inches(0.8),Inches(1.68),Inches(11.7),bh,s_dict['items'],size=15,gap=9)
        if has_note:
            tf=tb(s,Inches(0.8),Inches(5.7),Inches(11.7),Inches(0.7))
            p=tf.paragraphs[0];p.line_spacing=1.15;para(p,s_dict['note'],13,TEALD,TEALD)
        if has_stat:
            rect(s,Inches(0.8),Inches(5.55),Inches(11.7),Pt(1.2),LINE)
            stat_row(s,Inches(0.8),Inches(5.72),Inches(11.5),s_dict['stat'])
        footer(s,s_dict['foot'],pg)
    elif T=='split':
        bullets(s,Inches(0.8),Inches(1.68),Inches(7.0),Inches(4.9),s_dict['items'],size=14,gap=9)
        a=s_dict['aside']; dark=a.get('dark')
        ax,ay,aw,ah=Inches(8.05),Inches(1.68),Inches(4.48),Inches(4.85)
        rect(s,ax,ay,aw,ah,(INK if dark else CARDBG),line=(None if dark else CARDLN),lw=1.2,round_=True)
        tf=tb(s,ax+Inches(0.3),ay+Inches(0.28),aw-Inches(0.6),Inches(0.4))
        r=tf.paragraphs[0].add_run();r.text=a['title'];setf(r,13.5,(MINT if dark else TEALD),True)
        if a.get('stat'):
            tf=tb(s,ax+Inches(0.3),ay+Inches(0.9),aw-Inches(0.6),Inches(3.6))
            for i,(big,lbl) in enumerate(a['stat']):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(10)
                r=p.add_run();r.text=big;setf(r,20,(MINT if dark else TEALD),True)
                p2=tf.add_paragraph();p2.space_after=Pt(2)
                r=p2.add_run();r.text=lbl;setf(r,11,(RGBColor(0x9F,0xB0,0xAD) if dark else MUTE),True)
        else:
            it2=[(x[0],x[1],x[2] if len(x)>2 else '') for x in a['items']]
            tf=tb(s,ax+Inches(0.3),ay+Inches(0.9),aw-Inches(0.6),ah-Inches(1.2))
            for i,it in enumerate(it2):
                lvl,tx,cls=it; base=(WHITE if dark else CLSCOL.get(cls,BODY)); bc=(MINT if dark else (base if cls else INK))
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(8);p.line_spacing=1.12
                if lvl==0:
                    r=p.add_run();r.text="●  ";setf(r,8.5,(MINT if dark else TEAL),True)
                para(p,tx,12.5,base,bc)
        footer(s,s_dict['foot'],pg)
    elif T=='table':
        rows=[s_dict['headers']]+s_dict['rows']; nH=len(s_dict['headers'])
        colw=s_dict.get('colw')
        if not colw:
            colw=[Inches(11.7/nH)]*nH
        note=s_dict.get('note')
        th=Inches(1.7)
        gt=s.shapes.add_table(len(rows),nH,Inches(0.8),th,sum(colw,Emu(0)),Inches(0.5)*len(rows)).table
        gt.first_row=False;gt.horz_banding=False
        for ci,cw in enumerate(colw): gt.columns[ci].width=cw
        for ri,row in enumerate(rows):
            gt.rows[ri].height=Inches(0.52) if ri==0 else Inches(0.62)
            for ci,val in enumerate(row):
                cell=gt.cell(ri,ci);cell.margin_left=Inches(0.09);cell.margin_right=Inches(0.07)
                cell.margin_top=Inches(0.03);cell.margin_bottom=Inches(0.03);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                if ri==0: cell.fill.solid();cell.fill.fore_color.rgb=INK
                else: cell.fill.solid();cell.fill.fore_color.rgb=(WHITE if ri%2==1 else LGRAY)
                tf=cell.text_frame;tf.word_wrap=True;p=tf.paragraphs[0]
                p.alignment=PP_ALIGN.LEFT if ci==0 else PP_ALIGN.CENTER
                base=(WHITE if ri==0 else (INK if ci==0 else BODY))
                bc=(WHITE if ri==0 else TEALD)
                para(p,str(val),12 if ri else 12.5,base,bc)
                if ri==0 or ci==0:
                    for rr in p.runs: rr.font.bold=True
        if note:
            tf=tb(s,Inches(0.8),Inches(6.3),Inches(11.7),Inches(0.5))
            p=tf.paragraphs[0];p.line_spacing=1.15;para(p,note,12,TEALD,TEALD)
        footer(s,s_dict['foot'],pg)
    return

def render_special(s_dict,pg):
    T=s_dict['t']
    if T=='key':
        s=slide();rect(s,0,0,SW,SH,INK)
        tf=tb(s,Inches(0.9),Inches(0.75),Inches(11),Inches(0.35))
        r=tf.paragraphs[0].add_run();r.text=s_dict['eyebrow'].upper();setf(r,12,MINT,True)
        tf=tb(s,Inches(0.9),Inches(1.2),Inches(11.5),Inches(1.0))
        p=tf.paragraphs[0];r=p.add_run();r.text=s_dict['headline'];setf(r,29,RGBColor(0xF4,0xF7,0xF5),True)
        msgs=s_dict['msgs'];n=len(msgs);top=2.55;step=min(0.92,(6.7-top)/n)
        for i,(lab,desc) in enumerate(msgs):
            y=Inches(top+i*step)
            rect(s,Inches(0.9),y,Inches(11.5),Pt(1.0),RGBColor(0x2A,0x39,0x42))
            tf=tb(s,Inches(0.9),y+Inches(0.1),Inches(3.1),Inches(step*0.9),anchor=MSO_ANCHOR.MIDDLE)
            r=tf.paragraphs[0].add_run();r.text=lab;setf(r,15,MINT,True)
            tf=tb(s,Inches(4.15),y+Inches(0.1),Inches(8.3),Inches(step*0.9),anchor=MSO_ANCHOR.MIDDLE)
            p=tf.paragraphs[0];p.line_spacing=1.12;para(p,desc,13,RGBColor(0xD3,0xDB,0xD8),RGBColor(0xEA,0xF3,0xF1))
        return
    if T=='refs':
        s=slide();rect(s,0,0,SW,SH,PAPER)
        header(s,"참고문헌 · References",s_dict['title'],("PubMed","ink"))
        refs=s_dict['refs'];half=(len(refs)+1)//2
        for col,chunk in enumerate([refs[:half],refs[half:]]):
            x=Inches(0.8+col*6.0);tf=tb(s,x,Inches(1.68),Inches(5.75),Inches(5.0))
            for i,(txt,url) in enumerate(chunk):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.space_after=Pt(7);p.line_spacing=1.1
                para(p,txt,10,BODY,TEALD,link=url)
        footer(s,"제목 클릭 시 PubMed로 이동",pg)
        return

pg=0
for sd in NLC:
    pg+=1
    if sd['t'] in ('key','refs'): render_special(sd,pg)
    else: render(sd,pg)

out="/home/user/ppt-work/문헌고찰_NLC_RLS_LSB/01_NLC/NLC_발표_임상원장.pptx"
prs.save(out)
print("saved", out, "slides", len(prs.slides._sldIdLst))
