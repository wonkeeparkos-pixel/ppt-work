# -*- coding: utf-8 -*-
"""요추 경막외 차단 문헌고찰 — PPTX 생성기.

build_leb.py 의 슬라이드 정의(S)를 그대로 읽어 편집 가능한 PPTX로 변환한다.
- 본문의 <b>/<i>/<br> 인라인 마크업을 실제 굵게/기울임 run 으로 변환
- photo 슬라이드는 사진이 있으면 삽입, 없으면 출처가 적힌 점선 플레이스홀더
"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

from ppt_lib import (Deck, NAVY, TEAL, TEALD, INK, MUTE, LGRAY, CARD, WHITE,
                     GREEN, AMBER, RED, LINE, ICE, MINT)
import build_leb as C

CLS = {'red': RED, 'green': GREEN, 'accent': TEALD, 'muted': MUTE}
TAGC = {'red': RED, 'green': GREEN, 'amber': AMBER, 'ink': NAVY, '': TEAL}


def parse_rich(s):
    """<b>/<i>/<br> 를 (텍스트, bold, italic) 목록으로. 그 외 태그는 제거."""
    s = re.sub(r'<br\s*/?>', ' ', s)
    out, bold, ital, buf = [], False, False, []
    for tok in re.split(r'(</?[bi]>)', s):
        if tok == '<b>':
            out.append((''.join(buf), bold, ital)); buf = []; bold = True
        elif tok == '</b>':
            out.append((''.join(buf), bold, ital)); buf = []; bold = False
        elif tok == '<i>':
            out.append((''.join(buf), bold, ital)); buf = []; ital = True
        elif tok == '</i>':
            out.append((''.join(buf), bold, ital)); buf = []; ital = False
        else:
            buf.append(re.sub(r'<[^>]+>', '', tok))
    out.append((''.join(buf), bold, ital))
    return [(t, b, i) for t, b, i in out if t]


def plain(s):
    return re.sub(r'<[^>]+>', '', re.sub(r'<br\s*/?>', ' ', s))


class LEBDeck(Deck):
    def rbullets(self, s, x, y, w, h, items, size=14.5, gap=8):
        """items: (level, html, cls) — 인라인 굵게 유지."""
        tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, it in enumerate(items):
            lv = it[0]; html = it[1]
            base = CLS.get(it[2] if len(it) > 2 and it[2] else '', INK)
            sz = size if lv <= 0 else size * 0.94
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap); p.line_spacing = 1.1; p.alignment = PP_ALIGN.LEFT
            mark = "" if lv < 0 else ("●  " if lv == 0 else "–  ")
            indent = "" if lv <= 0 else ("      " * lv)
            if mark or indent:
                r = p.add_run(); r.text = indent + mark
                self._f(r, sz * 0.72 if lv == 0 else sz, base if base is not INK else (TEAL if lv == 0 else MUTE), True)
            for txt, b, ital in parse_rich(html):
                r = p.add_run(); r.text = txt
                self._f(r, sz, base, b or (lv < 0), ital)
        return tb

    def rtable(self, s, x, y, w, headers, rows, col_w, hl=(), size=11.5, row_h=Inches(0.46)):
        data = [headers] + rows
        n, nc = len(data), len(col_w)
        gt = s.shapes.add_table(n, nc, x, y, sum(col_w, Emu(0)), row_h * n).table
        gt.first_row = False; gt.horz_banding = False
        for ci, cw in enumerate(col_w):
            gt.columns[ci].width = cw
        for ri, row in enumerate(data):
            gt.rows[ri].height = Inches(0.46) if ri == 0 else row_h
            for ci, val in enumerate(row):
                cell = gt.cell(ri, ci)
                cell.margin_left = Inches(0.09); cell.margin_right = Inches(0.07)
                cell.margin_top = Inches(0.03); cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                if ri == 0:
                    cell.fill.fore_color.rgb = NAVY
                elif (ri - 1) in hl:
                    cell.fill.fore_color.rgb = RGBColor(0xFB, 0xEA, 0xE6)
                else:
                    cell.fill.fore_color.rgb = WHITE if ri % 2 == 1 else LGRAY
                tf = cell.text_frame; tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                col = WHITE if ri == 0 else (RED if (ri - 1) in hl and ci == 0 else INK)
                for txt, b, ital in parse_rich(val) or [(plain(val), False, False)]:
                    r = p.add_run(); r.text = txt
                    self._f(r, 12.0 if ri == 0 else size, col, b or ri == 0 or ci == 0, ital)
        return gt

    def photo_slide(self, spec, tag_color=TEAL):
        s = self._slide()
        self.header(s, plain(spec['title']), spec['eyebrow'],
                    spec['tag'][0] if isinstance(spec.get('tag'), tuple) else spec.get('tag'),
                    tag_color)
        ph = spec['photos']
        L, T = Inches(0.8), Inches(1.5)
        BW, BH = Inches(7.85), Inches(4.95)
        gapx = Inches(0.18)
        cw = BW if len(ph) == 1 else Emu(int((BW - gapx) / 2))
        for i, p in enumerate(ph):
            x = L + (cw + gapx) * i
            if p.get('path'):
                self.rect(s, x, T, cw, BH, RGBColor(0x0C, 0x10, 0x13), line=LINE)
                s.shapes.add_picture(p['path'], x + Inches(0.06), T + Inches(0.06),
                                     width=cw - Inches(0.12))
                self.text(s, x + Inches(0.12), T + BH - Inches(0.5), cw - Inches(0.24), Inches(0.42),
                          [[(plain(p['what']) + " — 출처 " + p['ref'], 9.5, RGBColor(0xC9, 0xD4, 0xD0), True)]],
                          anchor=MSO_ANCHOR.MIDDLE)
            else:
                box = self.rect(s, x, T, cw, BH, RGBColor(0xF4, 0xF7, 0xF5), line=RGBColor(0xBC, 0xCA, 0xC5), lw=1.6)
                try:
                    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                except Exception:
                    pass
                inner = cw - Inches(0.5)
                self.text(s, x + Inches(0.25), T + Inches(1.35), inner, Inches(0.3),
                          [[("실제 임상 사진 자리", 10, RGBColor(0x8A, 0x9A, 0x94), True)]],
                          align=PP_ALIGN.CENTER)
                self.text(s, x + Inches(0.25), T + Inches(1.75), inner, Inches(1.0),
                          [[(t, 13, RED if b else INK, True)] for t, b, _ in [(plain(p['what']), True, 0)]],
                          align=PP_ALIGN.CENTER, ls=1.15)
                self.text(s, x + Inches(0.25), T + Inches(2.85), inner, Inches(1.0),
                          [[("출처 후보 · " + p['ref'], 9.5, TEALD, True)]],
                          align=PP_ALIGN.CENTER, ls=1.2)
                self.text(s, x + Inches(0.25), T + Inches(3.95), inner, Inches(0.3),
                          [[(p['file'], 9, RGBColor(0x8A, 0x9A, 0x94), True)]], align=PP_ALIGN.CENTER)
        a = spec['aside']
        warn = spec.get('warn')
        self.card(s, Inches(8.95), Inches(1.5), Inches(3.68), BH,
                  RGBColor(0xFC, 0xF4, 0xF2) if warn else CARD,
                  line=RGBColor(0xE7, 0xCF, 0xC9) if warn else LINE)
        self.text(s, Inches(9.2), Inches(1.75), Inches(3.2), Inches(0.35),
                  [[(a['title'], 12.5, RED if warn else TEALD, True)]])
        self.rbullets(s, Inches(9.2), Inches(2.22), Inches(3.2), Inches(4.0),
                      a['items'], size=10.5, gap=6)
        self.footer(s, spec['foot'])
        return s


def tagof(spec):
    t = spec.get('tag')
    if not t:
        return None, TEAL
    if isinstance(t, tuple):
        return t[0], TAGC.get(t[1], TEAL)
    return t, TEAL


COLW = {  # 슬라이드 번호(1-based) → 열 너비
    4:  [Inches(3.6), Inches(3.0), Inches(5.25)],   # 진단별 근거 수준
    8:  [Inches(3.0), Inches(2.4), Inches(2.6), Inches(3.85)],  # 접근법별 효과
    13: [Inches(2.6), Inches(3.4), Inches(5.85)],   # 황색인대 결손
    16: [Inches(3.0), Inches(5.4), Inches(3.45)],   # LOR A계열
    19: [Inches(3.0), Inches(5.2), Inches(3.65)],   # LOR C계열
    21: [Inches(3.1), Inches(4.9), Inches(3.85)],   # LOR D계열
    29: [Inches(4.6), Inches(3.3), Inches(3.95)],   # 혈관내 검출력
    36: [Inches(2.5), Inches(3.3), Inches(3.5), Inches(2.55)],  # 감별표
    37: [Inches(2.9), Inches(4.4), Inches(4.55)],   # 합병증
}


def build():
    d = LEBDeck()
    for idx, sp in enumerate(C.S):
        T = sp['t']
        tag, tc = tagof(sp)
        if T == 'title':
            d.title_slide(sp['eyebrow'], [plain(sp['title'])], sp['sub'], sp['order'], sp['series'])
        elif T == 'key':
            d.key_slide(sp['eyebrow'], plain(sp['headline']),
                        [(plain(a), plain(b)) for a, b in sp['msgs']])
        elif T == 'bullets':
            s = d._slide(); d.header(s, plain(sp['title']), sp['eyebrow'], tag, tc)
            d.rbullets(s, Inches(0.8), Inches(1.5), Inches(11.85), Inches(5.2),
                       sp['items'], size=14.5, gap=8)
            d.footer(s, sp['foot'])
        elif T == 'split':
            s = d._slide(); d.header(s, plain(sp['title']), sp['eyebrow'], tag, tc)
            d.rbullets(s, Inches(0.8), Inches(1.5), Inches(7.15), Inches(5.1),
                       sp['items'], size=13.5, gap=7)
            a = sp['aside']; dark = a.get('dark')
            d.card(s, Inches(8.25), Inches(1.55), Inches(4.38), Inches(4.9),
                   NAVY if dark else CARD, line=None if dark else LINE)
            d.text(s, Inches(8.55), Inches(1.8), Inches(3.8), Inches(0.4),
                   [[(a['title'], 12.5, ICE if dark else TEALD, True)]])
            items = a['items']
            if dark:
                items = [(it[0], re.sub(r'</?b>', '', it[1]), it[2] if len(it) > 2 else '') for it in items]
            tb = d.rbullets(s, Inches(8.55), Inches(2.3), Inches(3.78), Inches(4.0),
                            items, size=11, gap=6)
            if dark:
                for p in tb.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.color.rgb == INK:
                            r.font.color.rgb = WHITE
            d.footer(s, sp['foot'])
        elif T == 'table':
            s = d._slide(); d.header(s, plain(sp['title']), sp['eyebrow'], tag, tc)
            cw = COLW.get(idx + 1)
            if not cw:
                n = len(sp['headers']); cw = [Inches(11.85 / n)] * n
            nrow = len(sp['rows']) + 1
            rh = Inches(min(0.62, max(0.36, 4.7 / nrow)))
            d.rtable(s, Inches(0.8), Inches(1.45), Inches(11.85), sp['headers'], sp['rows'],
                     cw, hl=set(sp.get('hlrows', [])), size=11 if nrow > 6 else 12, row_h=rh)
            if sp.get('note'):
                d.text(s, Inches(0.8), Inches(6.5), Inches(11.8), Inches(0.5),
                       [[(t, 10.5, TEALD, b) for t, b, _ in parse_rich(sp['note'])]], ls=1.15)
            d.footer(s, sp['foot'])
        elif T == 'photo':
            d.photo_slide(sp, tc)
        elif T == 'refs':
            d.refs_slide(sp['title'], [plain(t) for t, _ in sp['refs']])
    out = os.path.join(C.OUT_DIR, "LEB_발표.pptx")
    n = d.save(out)
    print(f"PPTX  → {out}  ({n}장)")
    return out


if __name__ == '__main__':
    build()
