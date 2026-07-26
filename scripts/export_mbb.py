# -*- coding: utf-8 -*-
"""
HTML 덱 → PPTX + PDF.

슬라이드에 손으로 그린 SVG 해부·투시 도해가 많아, python-pptx 도형으로 다시 그리면
도해가 통째로 사라진다. 그래서 각 슬라이드를 고해상도로 렌더해 16:9 슬라이드에
전면 배치하고, 슬라이드의 텍스트는 **발표자 노트**에 넣어 검색·복사가 되게 한다.
"""
import os
import re
import shutil
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Inches

import content_mbb as C

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "문헌고찰_Lumbar_MBB_Facet")
HTML = os.path.join(OUT, "요추_MBB_후관절차단_발표.html")
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"
SHOT_W = 1600          # CSS px — device_scale_factor로 2배 확대해 3200px로 저장
SCALE = 2


def strip_tags(s):
    s = re.sub(r"<br\s*/?>", " ", str(s))
    s = re.sub(r"<[^>]+>", "", s)
    return re.sub(r"\s+", " ", s).strip()


def slide_notes(sl, n):
    """슬라이드 원문 텍스트를 발표자 노트용으로 평문화."""
    L = [f"[{n:02d}] {strip_tags(sl.get('title', ''))}"]
    if sl.get("kick"):
        L.append(strip_tags(sl["kick"]))
    if sl.get("desc"):
        L.append(strip_tags(sl["desc"]))
    if sl.get("sub"):
        L.append(strip_tags(sl["sub"]))
    for it in sl.get("items", []):
        L.append(("  · " if it[0] >= 1 else "· ") + strip_tags(it[1]))
    a = sl.get("aside")
    if a:
        L.append(f"[{strip_tags(a.get('title',''))}]")
        for it in a.get("items", []):
            L.append("· " + strip_tags(it[1]))
        for st in a.get("stat", []):
            L.append(f"· {strip_tags(st[0])} — {strip_tags(st[1])}")
        if a.get("note"):
            L.append(strip_tags(a["note"]))
    for c in sl.get("cards", []):
        L.append(f"[{strip_tags(c['name'])}] {strip_tags(c.get('sl',''))}")
        for it in c.get("items", []):
            L.append("· " + strip_tags(it[1]))
        if c.get("cf"):
            L.append("· " + strip_tags(c["cf"]))
    if sl.get("headers"):
        hs = [strip_tags(h[0] if isinstance(h, tuple) else h) for h in sl["headers"]]
        L.append(" | ".join(hs))
        for r in sl.get("rows", []):
            L.append(" | ".join(strip_tags(c) for c in r))
    for p in sl.get("panels", []):
        L.append(f"[{strip_tags(p['pt'])}] {strip_tags(p.get('pl',''))}")
    for m in sl.get("msgs", []):
        L.append(f"· {strip_tags(m[0])} — {strip_tags(m[1])}")
    for r in sl.get("refs", []):
        L.append("· " + strip_tags(r))
    for k in ("caption", "note", "stat"):
        v = sl.get(k)
        if k == "stat" and v:
            L.append(" / ".join(f"{strip_tags(a)} {strip_tags(b)}" for a, b, *_ in v))
        elif v and k != "stat":
            L.append(strip_tags(v))
    if sl.get("foot"):
        L.append("출처: " + strip_tags(sl["foot"]))
    return "\n".join(L)


def shoot(tmpdir):
    paths = []
    with sync_playwright() as p:
        b = p.chromium.launch(executable_path=CHROME)
        pg = b.new_page(viewport={"width": SHOT_W, "height": int(SHOT_W * 9 / 16) + 60},
                        device_scale_factor=SCALE, color_scheme="light")
        pg.goto("file://" + HTML)
        pg.wait_for_timeout(2000)
        n = pg.evaluate("document.querySelectorAll('.snap').length")
        for i in range(n):
            pg.evaluate(f"document.querySelectorAll('.snap')[{i}].scrollIntoView()")
            pg.wait_for_timeout(200)
            f = os.path.join(tmpdir, f"s{i+1:02d}.png")
            pg.locator(".snap").nth(i).locator(".stage").screenshot(path=f)
            paths.append(f)
        b.close()
    return paths


def build_pptx(pngs, out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, png in enumerate(pngs):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(png, Emu(0), Emu(0),
                             width=prs.slide_width, height=prs.slide_height)
        if i < len(C.SLIDES):
            s.notes_slide.notes_text_frame.text = slide_notes(C.SLIDES[i], i + 1)
    prs.save(out)
    return out


def build_pdf(out):
    subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-sandbox",
                    "--no-pdf-header-footer", "--print-to-pdf-no-header",
                    f"--print-to-pdf={out}", "file://" + HTML],
                   check=True, capture_output=True, timeout=300)
    return out


if __name__ == "__main__":
    tmp = os.path.join(BASE, "_shots")
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp)
    pngs = shoot(tmp)
    print(f"슬라이드 {len(pngs)}장 렌더 ({SHOT_W*SCALE}px)")
    p = build_pptx(pngs, os.path.join(OUT, "요추_MBB_후관절차단_발표.pptx"))
    print("→", p, f"{os.path.getsize(p)/1e6:.1f}MB")
    d = build_pdf(os.path.join(OUT, "요추_MBB_후관절차단_발표.pdf"))
    print("→", d, f"{os.path.getsize(d)/1e6:.1f}MB")
    shutil.rmtree(tmp, ignore_errors=True)
