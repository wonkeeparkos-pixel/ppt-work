# -*- coding: utf-8 -*-
"""NLC 발표 — 동영상 없는 정적 PPTX(편집·저장용).
각 웹 슬라이드를 16:9 PNG로 래스터화해 풀블리드 배치. 임베드 미디어/타이밍 없음
→ 어느 PowerPoint에서든 '보호된 보기'만 해제하면 경고 없이 열리고 저장됨."""
import os, glob
from pptx import Presentation
from pptx.util import Inches

BASE = "/home/user/ppt-work/문헌고찰_NLC_RLS_LSB/01_NLC"
SCRATCH = "/tmp/claude-0/-home-user-ppt-work/bb19d1cb-d1ba-542e-8235-38fcac774773/scratchpad"
PNG_DIR = os.path.join(SCRATCH, "png")
OUT     = os.path.join(BASE, "NLC_발표_v2.7.pptx")

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

pngs = sorted(glob.glob(os.path.join(PNG_DIR, "s*.png")))
assert len(pngs) == 22, f"expected 22 pngs, got {len(pngs)}"

for png in pngs:
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.save(OUT)
print("saved:", OUT, "|", round(os.path.getsize(OUT)/1e6, 2), "MB | slides", len(prs.slides._sldIdLst))
