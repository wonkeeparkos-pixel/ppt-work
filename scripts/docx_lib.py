# -*- coding: utf-8 -*-
"""Markdown → DOCX 렌더러 (한국어 의학 문헌고찰용). 같은 마크다운을 Google 문서 업로드에도 재사용."""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

KFONT = "맑은 고딕"
NAVY = RGBColor(0x1B, 0x35, 0x5E)
TEAL = RGBColor(0x1F, 0x6E, 0x70)
MUTE = RGBColor(0x55, 0x60, 0x6B)
RED  = RGBColor(0xA3, 0x30, 0x25)
INK  = RGBColor(0x1A, 0x20, 0x2C)

def _kfont(run, size=None, color=None, bold=None, italic=None):
    run.font.name = 'Calibri'
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), KFONT); rf.set(qn('w:ascii'), 'Calibri'); rf.set(qn('w:hAnsi'), 'Calibri')
    if size is not None: run.font.size = Pt(size)
    if color is not None: run.font.color.rgb = color
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic

def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),hexcolor); tcPr.append(sh)

def _cant_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))
    if header: trPr.append(OxmlElement('w:tblHeader'))

def _add_runs(p, text, base_size=10.5, base_color=INK, base_bold=False):
    """**bold** 인라인 처리."""
    for i, seg in enumerate(re.split(r'(\*\*.+?\*\*)', text)):
        if not seg: continue
        if seg.startswith('**') and seg.endswith('**'):
            r = p.add_run(seg[2:-2]); _kfont(r, base_size, base_color, True)
        else:
            r = p.add_run(seg); _kfont(r, base_size, base_color, base_bold)

def render_markdown(md, out_path):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), KFONT)
    for sec in doc.sections:
        sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
        sec.left_margin = Inches(0.95); sec.right_margin = Inches(0.95)

    lines = md.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        # table
        if s.startswith('|') and i+1 < len(lines) and set(lines[i+1].strip().replace('|','').replace(':','').replace('-','').replace(' ','')) == set():
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells); i += 1
            header = rows[0]; body = rows[2:]  # skip separator row
            t = doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.style = 'Table Grid'; _cant_split(t.rows[0], header=True)
            for ci, h in enumerate(header):
                c = t.rows[0].cells[ci]; _shade(c, "1B355E")
                c.vertical_anchor = 1
                pp = c.paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = pp.add_run(h); _kfont(rr, 9.3, RGBColor(0xFF,0xFF,0xFF), True)
            for ri, row in enumerate(body):
                cells = t.add_row(); _cant_split(cells)
                for ci, val in enumerate(row):
                    cell = cells.cells[ci]
                    if ri % 2 == 0: _shade(cell, "EEF2F6")
                    pp = cell.paragraphs[0]
                    pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    pp.paragraph_format.space_before = Pt(1); pp.paragraph_format.space_after = Pt(1)
                    _add_runs(pp, val, 9.2, INK)
            for ci in range(len(header)):
                pass
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if s.startswith('# '):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(3)
            r = p.add_run(s[2:]); _kfont(r, 19, NAVY, True)
        elif s.startswith('## '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(13); p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'10'); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),'1F6E70')
            pbdr.append(bottom); pPr.append(pbdr)
            r = p.add_run(s[3:]); _kfont(r, 14.5, NAVY, True)
        elif s.startswith('### '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(9); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            rr = p.add_run("▍ "); _kfont(rr, 11, TEAL, True)
            r = p.add_run(s[4:]); _kfont(r, 11.5, NAVY, True)
        elif s.startswith('> '):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
            _add_runs(p, s[2:], 9.6, MUTE);
            for rn in p.runs: rn.font.italic = True
        elif re.match(r'^\d+\.\s', s):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.08
            num = re.match(r'^(\d+\.)\s', s).group(1)
            rn = p.add_run(num + " "); _kfont(rn, 10.3, TEAL, True)
            _add_runs(p, s[len(num):].strip(), 10.3, INK)
        elif s.startswith('- ') or line.startswith('  - '):
            lvl = 1 if line.startswith('  - ') else 0
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3 + lvl*0.28)
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.08
            mark = "•  " if lvl == 0 else "–  "
            rn = p.add_run(mark); _kfont(rn, 9 if lvl==0 else 10.3, TEAL if lvl==0 else MUTE, True)
            _add_runs(p, s[2:], 10.3, INK)
        elif s == '---':
            pass
        elif s == '':
            pass
        else:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.12
            _add_runs(p, s, 10.5, INK)
        i += 1
    doc.save(out_path)
    return out_path
