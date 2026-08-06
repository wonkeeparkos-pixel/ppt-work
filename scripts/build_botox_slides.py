# -*- coding: utf-8 -*-
"""구글 슬라이드용 **가벼운 편집본** PPTX.

디자인 덱(`보툴리눔_정형외과통증_발표.pptx`)은 각 쪽이 고해상도 이미지라 9.5 MB다.
그 파일은 구글 드라이브 API로 올릴 수 없어서(내용을 통째로 전송해야 한다),
같은 내용을 **글자와 도형으로만** 다시 짜서 200 KB 이하로 만든다.

이 파일은 구글 드라이브에 올리면 구글 슬라이드로 변환돼 **브라우저에서 바로 고칠 수 있다.**
발표용 최종본은 어디까지나 디자인 덱 쪽이다 — 이건 내용 확인·수정용 사본이다.

    python3 scripts/build_botox_slides.py
"""
import os
import re
import sys
import html as _html

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from content_botox import BOTOX  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "문헌고찰_보툴리눔_정형외과통증", "보툴리눔_정형외과통증_슬라이드편집본.pptx")

# 차단 / Blockade 팔레트 — 웹 덱과 같은 값
PAPER = RGBColor(0xF4, 0xF3, 0xF8)
INK = RGBColor(0x18, 0x1A, 0x2E)
INK2 = RGBColor(0x33, 0x35, 0x5A)
INDIGO = RGBColor(0x46, 0x3A, 0x96)
BRICK = RGBColor(0xA8, 0x35, 0x2A)
MOSS = RGBColor(0x2E, 0x6B, 0x3E)
AMBER = RGBColor(0xB0, 0x70, 0x1A)
SLATE = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TAGCOLOR = {"green": MOSS, "amber": AMBER, "red": BRICK,
            "ink": INK, "slate": SLATE, "": INDIGO}
FONT = "Noto Sans KR"          # 구글 슬라이드 기본 탑재 한글 폰트

W, H = Inches(13.333), Inches(7.5)
L = Inches(0.62)               # 좌측 여백
CW = W - 2 * L                 # 본문 폭


def txt(s):
    """HTML 조각을 평문으로."""
    if s is None:
        return ""
    s = re.sub(r"<br\s*/?>", "\n", str(s))
    s = re.sub(r"<[^>]+>", "", s)
    return _html.unescape(s).strip()


def tag_of(s):
    """('라벨','색') 또는 '라벨' → (라벨, RGB)."""
    t = s.get("tag")
    if not t:
        return None, INDIGO
    if isinstance(t, (tuple, list)):
        return txt(t[0]), TAGCOLOR.get(t[1], INDIGO)
    return txt(t), INDIGO


def box(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def line(tf, text, size, color, bold=False, space_before=0, first=False,
         align=PP_ALIGN.LEFT, spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    if space_before:
        p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    f = r.font
    f.size, f.bold, f.name = Pt(size), bold, FONT
    f.color.rgb = color
    return p


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def chrome(slide, s, dark=False):
    """상단 등급바 · eyebrow · 태그칩 · 제목 · 각주. 본문 시작 y를 돌려준다."""
    label, color = tag_of(s)
    rect(slide, 0, 0, W, Pt(4.5), color)                      # 등급 = 색

    if s.get("eyebrow"):
        tf = box(slide, L, Inches(0.30), CW * 0.68, Inches(0.3))
        line(tf, txt(s["eyebrow"]), 10.5, WHITE if dark else INDIGO,
             bold=True, first=True)

    if label:
        cw = Inches(0.13) * max(len(label), 4) + Inches(0.34)
        ch = rect(slide, W - L - cw, Inches(0.26), cw, Inches(0.30),
                  color, MSO_SHAPE.ROUNDED_RECTANGLE)
        ch.adjustments[0] = 0.28
        tf = ch.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        line(tf, label, 9.5, WHITE, bold=True, first=True, align=PP_ALIGN.CENTER)

    y = Inches(0.72)
    if s.get("title"):
        t = txt(s["title"])
        size = 27 if len(t) <= 34 else (24 if len(t) <= 48 else 21)
        tf = box(slide, L, y, CW, Inches(1.0))
        line(tf, t, size, WHITE if dark else INK, bold=True, first=True, spacing=1.08)
        y += Inches(0.62) + Inches(0.30) * (len(t) // 40)

    if s.get("foot"):
        tf = box(slide, L, H - Inches(0.62), CW, Inches(0.4))
        line(tf, txt(s["foot"]), 8.5, SLATE, first=True, spacing=1.15)
    return y + Inches(0.30)


def bullets(tf, items, size=12.5, first=True):
    """items = [(level, text, cls), ...]"""
    for it in items:
        lvl, body = it[0], it[1]
        cls = it[2] if len(it) > 2 else ""
        color = {"green": MOSS, "red": BRICK}.get(cls, INK2)
        prefix = "→ " if lvl == -1 else "· "
        p = line(tf, prefix + txt(body), size, color,
                 bold=(lvl == -1), space_before=0 if first else 7,
                 first=first, spacing=1.22)
        if lvl > 0:
            p.level = min(lvl, 4)
        first = False


def stats(slide, y, data, dark=False):
    """(숫자, 라벨) 카드 줄."""
    n = len(data)
    gap = Inches(0.16)
    w = (CW - gap * (n - 1)) / n
    for i, (num, lab) in enumerate(data):
        x = L + (w + gap) * i
        rect(slide, x, y, w, Inches(0.92), INDIGO if dark else WHITE)
        tf = box(slide, x + Inches(0.14), y + Inches(0.10), w - Inches(0.28), Inches(0.72))
        line(tf, txt(num), 19, WHITE if dark else INDIGO, bold=True, first=True)
        line(tf, txt(lab), 8.5, WHITE if dark else SLATE, spacing=1.1)
    return y + Inches(1.06)


def add_table(slide, s, y):
    headers = [txt(h) for h in s["headers"]]
    rows = s["rows"]
    n_r, n_c = len(rows) + 1, len(headers)
    h = min(Inches(4.9), Inches(0.34) * n_r)
    gt = slide.shapes.add_table(n_r, n_c, L, y, CW, h).table
    hl = set(s.get("hlrows") or [])
    fs = 9 if s.get("dense") or n_r > 9 else 10.5

    for c, head in enumerate(headers):
        cell = gt.cell(0, c)
        cell.text = head
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.name = Pt(fs), True, FONT
                r.font.color.rgb = WHITE

    for r_i, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = gt.cell(r_i, c)
            cell.text = txt(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                RGBColor(0xE8, 0xE6, 0xF4) if (r_i - 1) in hl else
                (PAPER if r_i % 2 else WHITE))
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.name = Pt(fs), FONT
                    r.font.color.rgb = INK2
    y += h + Inches(0.12)
    if s.get("note"):
        tf = box(slide, L, y, CW, Inches(0.5))
        line(tf, txt(s["note"]), 8.5, INDIGO, first=True, spacing=1.2)
    return y


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    for idx, s in enumerate(BOTOX, 1):
        sl = prs.slides.add_slide(blank)
        t = s["t"]
        dark = t in ("title", "key")
        rect(sl, 0, 0, W, H, INK if dark else PAPER)

        # ── 표지 ──────────────────────────────────────────
        if t == "title":
            _, color = tag_of(s)
            rect(sl, 0, 0, W, Pt(4.5), INDIGO)
            tf = box(sl, L, Inches(1.5), CW, Inches(0.4))
            line(tf, txt(s["eyebrow"]), 12, RGBColor(0x8B, 0x7B, 0xE8), bold=True, first=True)
            tf = box(sl, L, Inches(2.15), CW, Inches(2.2))
            line(tf, txt(s["title"]), 48, WHITE, bold=True, first=True, spacing=1.02)
            tf = box(sl, L, Inches(4.45), CW, Inches(1.4))
            line(tf, txt(s["sub"]), 15, RGBColor(0xC9, 0xC4, 0xE6), first=True, spacing=1.3)
            line(tf, txt(s["order"]), 11.5, RGBColor(0x8B, 0x7B, 0xE8), bold=True, space_before=12)
            tf = box(sl, L, H - Inches(0.75), CW, Inches(0.4))
            line(tf, txt(s["series"]), 9, SLATE, first=True)
            continue

        y = chrome(sl, s, dark)

        # ── 핵심 요약(짙은 배경) ──────────────────────────
        if t == "key":
            if s.get("headline"):
                tf = box(sl, L, Inches(0.72), CW, Inches(0.8))
                line(tf, txt(s["headline"]), 26, WHITE, bold=True, first=True, spacing=1.08)
            y = Inches(1.75)
            for head, body in s["msgs"]:
                tf = box(sl, L, y, CW, Inches(0.9))
                line(tf, txt(head), 12, RGBColor(0x8B, 0x7B, 0xE8), bold=True, first=True)
                line(tf, txt(body), 11.5, RGBColor(0xDD, 0xDA, 0xEE), space_before=2, spacing=1.25)
                y += Inches(0.92)
            continue

        # ── 참고문헌 2단 ──────────────────────────────────
        if t == "refs":
            refs = s["refs"]
            half = (len(refs) + 1) // 2
            for col, chunk in enumerate((refs[:half], refs[half:])):
                if not chunk:
                    continue
                x = L + (CW / 2 + Inches(0.12)) * col
                tf = box(sl, x, y, CW / 2 - Inches(0.12), Inches(5.2))
                for i, ref in enumerate(chunk):
                    line(tf, txt(ref[0]), 10, INK, bold=True,
                         space_before=0 if i == 0 else 9, first=(i == 0), spacing=1.18)
                    if len(ref) > 1 and ref[1]:
                        line(tf, txt(ref[1]), 8.5, SLATE, spacing=1.15)
            continue

        # ── 표 ────────────────────────────────────────────
        if t == "table":
            add_table(sl, s, y)
            continue

        # ── 일러스트 슬라이드 ─────────────────────────────
        if t in ("figure", "bigfig"):
            if s.get("caption"):
                tf = box(sl, L, y, CW, Inches(1.2))
                line(tf, txt(s["caption"]), 13, INK2, first=True, spacing=1.3)
                y += Inches(0.95)
            bx = rect(sl, L, y, CW, H - y - Inches(0.95), WHITE)
            tf = bx.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            line(tf, "［ 일러스트 ］", 13, INDIGO, bold=True, first=True, align=PP_ALIGN.CENTER)
            line(tf, "이 쪽의 그림은 디자인 덱(웹 HTML · PPTX)에 있습니다.\n"
                     "이 편집본은 글자만 담은 사본입니다.",
                 10, SLATE, space_before=6, align=PP_ALIGN.CENTER, spacing=1.25)
            continue

        # ── 불릿 / 좌우분할 ───────────────────────────────
        body_h = H - y - Inches(0.85)
        if t == "split" and s.get("aside"):
            bw = CW * 0.575
            aw = CW - bw - Inches(0.22)
            tf = box(sl, L, y, bw, body_h)
            bullets(tf, s.get("items", []))

            a = s["aside"]
            ax = L + bw + Inches(0.22)
            a_dark = bool(a.get("dark"))
            rect(sl, ax, y, aw, body_h, INK if a_dark else WHITE)
            tf = box(sl, ax + Inches(0.18), y + Inches(0.16), aw - Inches(0.36),
                     body_h - Inches(0.32))
            first = True
            if a.get("title"):
                line(tf, txt(a["title"]), 11.5, WHITE if a_dark else INDIGO,
                     bold=True, first=True)
                first = False
            if a.get("items"):
                for it in a["items"]:
                    lvl, bodytxt = it[0], it[1]
                    cls = it[2] if len(it) > 2 else ""
                    color = {"green": MOSS, "red": BRICK}.get(
                        cls, RGBColor(0xDD, 0xDA, 0xEE) if a_dark else INK2)
                    line(tf, ("→ " if lvl == -1 else "· ") + txt(bodytxt), 10.5, color,
                         bold=(lvl == -1), space_before=0 if first else 7,
                         first=first, spacing=1.22)
                    first = False
            if a.get("stat"):
                for num, lab in a["stat"]:
                    line(tf, txt(num), 17, WHITE if a_dark else INDIGO, bold=True,
                         space_before=0 if first else 8, first=first)
                    line(tf, txt(lab), 8.5, SLATE, spacing=1.15)
                    first = False
        else:
            tf = box(sl, L, y, CW, body_h)
            bullets(tf, s.get("items", []))

        if s.get("stat") and t != "split":
            stats(sl, H - Inches(1.72), s["stat"])
        if s.get("note"):
            tf = box(sl, L, H - Inches(0.92), CW, Inches(0.35))
            line(tf, txt(s["note"]), 8.5, INDIGO, first=True, spacing=1.2)

    prs.save(OUT)
    kb = os.path.getsize(OUT) / 1024
    print(f"slides: {len(prs.slides.__iter__.__self__._sldIdLst)} | {kb:.0f} KB -> {OUT}")


if __name__ == "__main__":
    build()
