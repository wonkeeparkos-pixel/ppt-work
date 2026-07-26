# -*- coding: utf-8 -*-
"""손 경련 웹덱 → 정적 PPTX (NLC 정적 PPTX와 동일 방식).
각 슬라이드(.stage)를 16:9 PNG로 래스터화해 풀블리드 배치. 임베드 미디어 없음."""
import os, glob, sys
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

BASE = "/home/user/ppt-work/문헌고찰_손경련"
HTML = os.path.join(BASE, "손경련_발표_웹_v1.0.html")
OUT = os.path.join(BASE, "손경련_발표_v1.0.pptx")
PNG_DIR = "/tmp/claude-0/-home-user-ppt-work/c743b03e-5647-5eeb-b648-75f6a5c6d985/scratchpad/hand_png"

os.makedirs(PNG_DIR, exist_ok=True)
for f in glob.glob(os.path.join(PNG_DIR, "*.png")):
    os.remove(f)

with sync_playwright() as p:
    br = p.chromium.launch(executable_path="/opt/pw-browsers/chromium-1194/chrome-linux/chrome")
    pg = br.new_page(viewport={"width": 1600, "height": 900}, device_scale_factor=2)
    pg.goto("file://" + HTML)
    pg.wait_for_timeout(2500)  # 폰트 로드
    pg.add_style_tag(content=".a-cramp,.a-spark,.a-sway,.a-brk,.a-flow{animation:none!important}"
                             ".bar,.count,.hint,.deckttl{display:none!important}")
    stages = pg.query_selector_all(".stage")
    print("stages:", len(stages))
    for i, st in enumerate(stages, 1):
        st.scroll_into_view_if_needed()
        pg.wait_for_timeout(120)
        st.screenshot(path=os.path.join(PNG_DIR, f"s{i:02d}.png"))
    pg.pdf(path=os.path.join(BASE, "손경련_발표_웹_미리보기.pdf"),
           width="13.333in", height="7.5in", print_background=True, margin={"top": "0", "bottom": "0"})
    br.close()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
pngs = sorted(glob.glob(os.path.join(PNG_DIR, "s*.png")))
for png in pngs:
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
prs.save(OUT)
print("saved:", OUT, "|", round(os.path.getsize(OUT) / 1e6, 2), "MB | slides", len(prs.slides._sldIdLst))
