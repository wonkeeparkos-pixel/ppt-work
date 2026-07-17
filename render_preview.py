# -*- coding: utf-8 -*-
"""생성된 pptx를 실제 파싱하여 PNG로 렌더링 (레이아웃 검증용 미리보기)."""
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

DPI = 96.0
SCALE = DPI / 914400.0        # EMU -> px
PT2PX = DPI / 72.0            # point -> px (핵심: 실제 크기 반영)
FONTPATH = "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"
_fc = {}
def font(size_pt, bold=False):
    px = max(6, int(round(size_pt * PT2PX)))
    k = (px, bold)
    if k not in _fc:
        _fc[k] = ImageFont.truetype(FONTPATH, px)
    return _fc[k]

def emu_px(v): return int(round(v * SCALE))

def rgb(c):
    try: return (c[0], c[1], c[2])
    except Exception: return None

def get_fill(sh):
    try:
        if sh.fill.type == 1:  # SOLID
            return rgb(sh.fill.fore_color.rgb)
    except Exception: pass
    return None

def get_line(sh):
    try:
        ln = sh.line
        if ln.fill.type == 1:
            w = ln.width.pt if ln.width else 1.0
            return rgb(ln.color.rgb), max(1, int(round(w)))
    except Exception: pass
    return None, 0

def run_color(r):
    try:
        if r.font.color and r.font.color.type is not None:
            return rgb(r.font.color.rgb)
    except Exception: pass
    return (30, 30, 30)

def draw_text_frame(draw, sh, x, y, w, h):
    tf = sh.text_frame
    # vertical anchor
    va = tf.vertical_anchor  # None/TOP=1/MIDDLE=3/BOTTOM=4
    pad = 4
    ix, iw = x + pad, w - 2 * pad
    # build lines
    paras = []
    for p in tf.paragraphs:
        runs = [(r.text, (r.font.size.pt if r.font.size else 18),
                 bool(r.font.bold), run_color(r)) for r in p.runs if r.text != ""]
        align = p.alignment  # 1 L,2 C,3 R? PP_ALIGN: LEFT=1,CENTER=2,RIGHT=3
        ls = p.line_spacing if p.line_spacing else 1.0
        if not runs:
            paras.append(([], 12 * PT2PX * (ls if isinstance(ls, float) else 1.0), align, ls))
            continue
        # char-level flow -> list of lines; each line = list of (char,size,bold,color)
        chars = []
        for t, sz, b, col in runs:
            for ch in t:
                chars.append((ch, sz, b, col))
        lines, cur, curw = [], [], 0
        for ch, sz, b, col in chars:
            cw = font(sz, b).getlength(ch)
            if ch == "\n" or (cur and curw + cw > iw):
                if ch == "\n":
                    lines.append(cur); cur, curw = [], 0; continue
                lines.append(cur); cur, curw = [], 0
            cur.append((ch, sz, b, col)); curw += cw
        if cur: lines.append(cur)
        maxsz = max((c[1] for ln in lines for c in ln), default=14)
        lh = maxsz * PT2PX * 1.32 * (ls if isinstance(ls, float) else 1.0)
        for ln in lines:
            paras.append((ln, lh, align, ls))
    total_h = sum(lh for _, lh, _, _ in paras)
    if va == 3:   ty = y + (h - total_h) / 2
    elif va == 4: ty = y + (h - total_h) - pad
    else:         ty = y + pad
    cy = ty
    for ln, lh, align, ls in paras:
        lw = sum(font(sz, b).getlength(ch) for ch, sz, b, col in ln)
        if align == 2:   lx = ix + (iw - lw) / 2
        elif align == 3: lx = ix + (iw - lw)
        else:            lx = ix
        cx = lx
        for ch, sz, b, col in ln:
            f = font(sz, b)
            draw.text((cx, cy), ch, font=f, fill=col or (30, 30, 30))
            if b: draw.text((cx + 1, cy), ch, font=f, fill=col or (30, 30, 30))
            cx += f.getlength(ch)
        cy += lh

def render(pptx, outdir):
    prs = Presentation(pptx)
    W = emu_px(prs.slide_width); H = emu_px(prs.slide_height)
    imgs = []
    for i, slide in enumerate(prs.slides, 1):
        # background
        bg = (255, 255, 255)
        try:
            if slide.background.fill.type == 1:
                bg = rgb(slide.background.fill.fore_color.rgb) or bg
        except Exception: pass
        img = Image.new("RGB", (W, H), bg)
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            x, y = emu_px(sh.left or 0), emu_px(sh.top or 0)
            w, hh = emu_px(sh.width or 0), emu_px(sh.height or 0)
            fill = get_fill(sh)
            lc, lw = get_line(sh)
            st = None
            try: st = sh.auto_shape_type
            except Exception: st = None
            if fill or lc:
                if st is not None and int(st) == 9:      # OVAL
                    d.ellipse([x, y, x + w, y + hh], fill=fill, outline=lc, width=lw or 1)
                elif st is not None and int(st) == 5:    # ROUNDED_RECT
                    r = max(2, int(min(w, hh) * 0.12))
                    d.rounded_rectangle([x, y, x + w, y + hh], radius=r, fill=fill, outline=lc, width=lw or 1)
                else:
                    d.rectangle([x, y, x + w, y + hh], fill=fill, outline=lc, width=lw or 1)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_text_frame(d, sh, x, y, w, hh)
        p = f"{outdir}/slide_{i:02d}.png"
        img.save(p); imgs.append(img)
    # contact sheet 5 x 6
    cols, rows = 5, 6
    tw, th = W // 4, H // 4
    sheet = Image.new("RGB", (cols * tw + (cols + 1) * 10, rows * th + (rows + 1) * 10), (235, 238, 237))
    for idx, im in enumerate(imgs):
        r, c = divmod(idx, cols)
        t = im.resize((tw, th))
        sheet.paste(t, (10 + c * (tw + 10), 10 + r * (th + 10)))
    sheet.save(f"{outdir}/_contact_sheet.png")
    print("rendered", len(imgs), "slides ->", outdir)

if __name__ == "__main__":
    import sys
    render(sys.argv[1], sys.argv[2])
