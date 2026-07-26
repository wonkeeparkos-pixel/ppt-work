# -*- coding: utf-8 -*-
"""시술실 벽 부착용 A4 요약 1장 (세로)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

INK   = RGBColor(0x1B, 0x2C, 0x38)
DEEP  = RGBColor(0x12, 0x4E, 0x66)
TEAL  = RGBColor(0x1B, 0x87, 0x9E)
MINT  = RGBColor(0xEA, 0xF4, 0xF6)
MINT2 = RGBColor(0xD5, 0xE9, 0xED)
CORAL = RGBColor(0xC4, 0x54, 0x4A)
CORALB= RGBColor(0xFA, 0xEB, 0xE9)
GOLD  = RGBColor(0xD9, 0x8E, 0x2B)
GRAY  = RGBColor(0x74, 0x86, 0x8E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Malgun Gothic"

prs = Presentation()
prs.slide_width  = Inches(8.27)   # A4 세로
prs.slide_height = Inches(11.69)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = PAPER

def _ea(r, n):
    rPr = r._r.get_or_add_rPr()
    for t in ("a:ea","a:cs","a:latin"):
        el = rPr.find(qn(t))
        if el is None: el = rPr.makeelement(qn(t), {}); rPr.append(el)
        el.set("typeface", n)

def rect(x,y,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,radius=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp

def text(x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.0):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    paras = [[runs]] if isinstance(runs, tuple) else [[p] if isinstance(p,tuple) else list(p) for p in runs]
    for i,para in enumerate(paras):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = ls
        for seg in para:
            r = p.add_run(); r.text = seg[0]
            r.font.size = Pt(seg[1]); r.font.bold = seg[3]
            r.font.color.rgb = seg[2]; r.font.name = FONT
            _ea(r, FONT)
    return tb

# 헤더
rect(0, 0, 8.27, 1.15, fill=DEEP)
rect(0, 1.15, 8.27, 0.08, fill=GOLD)
text(0.55, 0.24, 7.2, 0.5, ("시술 멘트 요약", 24, WHITE, True))
text(0.57, 0.76, 7.2, 0.32, ("아픈 걸 숨기지 말고, 끝을 알려주세요", 13, RGBColor(0x9A,0xD1,0xDB), False))

# 순서대로 말하기
y = 1.42
text(0.55, y, 7.2, 0.32, ("순서대로 이렇게", 14, DEEP, True))
y += 0.38
steps = [
    ("찌르기 전", "잠깐 따끔하고 몇 초면 지나가요. 셋 셀게요 — 하나, 둘, 셋.", TEAL),
    ("바늘 진행", "이제부턴 지그시 누르는 느낌이에요. 날카로우면 말씀하세요.", DEEP),
    ("시술 중", "절반쯤 지났어요. 잘하고 계세요. 숨 편하게 쉬세요.", TEAL),
    ("끝난 직후", "잘 참으셨어요. 어지럽거나 불편한 데 없으세요?", DEEP),
    ("보내기 전", "내일 잠깐 뻐근할 수 있어요. 이상하면 바로 전화 주세요.", GOLD),
]
for i,(when, line, col) in enumerate(steps):
    rect(0.55, y, 1.55, 0.56, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    text(0.55, y, 1.55, 0.56, (when, 11.5, WHITE, True), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(2.22, y, 5.5, 0.56, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    text(2.5, y, 5.1, 0.56, [[("“"+line+"”", 11.5, INK, False)]], anchor=MSO_ANCHOR.MIDDLE, ls=1.2)
    y += 0.62

# 감각 단어
y += 0.12
text(0.55, y, 7.2, 0.32, ("감각은 단계마다 다르게", 14, DEEP, True))
y += 0.4
feels = [("피부 뚫을 때","따끔"),("마취약","시큰·화끈"),("바늘 진행","누르는 느낌"),("약 주입","뻐근·차오름")]
for i,(a,b) in enumerate(feels):
    x = 0.55 + i*1.86
    rect(x, y, 1.72, 0.75, fill=MINT2, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(x, y+0.11, 1.72, 0.28, (a, 10, GRAY, False), align=PP_ALIGN.CENTER)
    text(x, y+0.37, 1.72, 0.32, (b, 13, DEEP, True), align=PP_ALIGN.CENTER)
y += 0.88

# 금지어
text(0.55, y, 7.2, 0.32, ("이 말은 빼기", 14, CORAL, True))
y += 0.38
rect(0.55, y, 7.17, 1.2, fill=CORALB, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
bads = ["“여기가 제일 아픈 부분이에요”", "“하나도 안 아파요”",
        "“많이 아팠죠?” (끝난 뒤)", "“절대 움직이면 안 돼요”"]
for i,b in enumerate(bads):
    x = 0.85 + (i%2)*3.4
    yy = y + 0.18 + (i//2)*0.46
    text(x, yy, 3.2, 0.4, [[("× ", 12, CORAL, True), (b, 11.5, INK, False)]])
y += 1.36

# 곤란한 순간
text(0.55, y, 7.2, 0.32, ("곤란할 때 첫 문장", 14, DEEP, True))
y += 0.38
scenes = [
    ("어지럽다고 할 때", "괜찮습니다, 바로 눕혀드릴게요. 다리 올릴게요."),
    ("바늘 무서워할 때", "무서워하시는 분 정말 많아요. 준비되면 말씀하세요."),
    ("다음 날 더 아플 때", "말씀드렸던 반응이에요. 2~3일이면 가라앉습니다."),
    ("효과 없다고 할 때", "언제부터 다시 아프기 시작했는지 알려주시겠어요?"),
]
for i,(a,b) in enumerate(scenes):
    rect(0.55, y, 7.17, 0.55, fill=MINT if i%2==0 else WHITE,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    if i%2==1:
        rect(0.55, y, 7.17, 0.55, fill=None, line=MINT2, lw=1.0,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.14)
    text(0.85, y, 1.95, 0.55, (a, 11, DEEP, True), anchor=MSO_ANCHOR.MIDDLE)
    text(2.9, y, 4.6, 0.55, [[("“"+b+"”", 11, INK, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.6

# 푸터
rect(0, 11.28, 8.27, 0.41, fill=DEEP)
text(0, 11.28, 8.27, 0.41, ("환자는 아프지 않아서가 아니라, 아플 때 곁에 있어줘서 다시 옵니다", 11, WHITE, False),
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save("/home/user/ppt-work/시술멘트_벽부착_A4.pptx")
print("wallcard OK")
