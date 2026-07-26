# -*- coding: utf-8 -*-
"""
의사용 멘트 매뉴얼 — 말로 통증을 줄이고 신뢰를 얻는 법
한국 의원 진료실에서 원장이 실제로 쓰는 문장 중심. 16:9.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 팔레트 ----------
INK   = RGBColor(0x1B, 0x2C, 0x38)
DEEP  = RGBColor(0x12, 0x4E, 0x66)   # 딥 블루틸
TEAL  = RGBColor(0x1B, 0x87, 0x9E)   # 포인트
TEALL = RGBColor(0x9A, 0xD1, 0xDB)
MINT  = RGBColor(0xEA, 0xF4, 0xF6)
MINT2 = RGBColor(0xD5, 0xE9, 0xED)
CORAL = RGBColor(0xC4, 0x54, 0x4A)   # 피할 말
CORALB= RGBColor(0xFA, 0xEB, 0xE9)
GOLD  = RGBColor(0xD9, 0x8E, 0x2B)
GRAY  = RGBColor(0x74, 0x86, 0x8E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xFB, 0xFD, 0xFD)

FONT = "Malgun Gothic"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def _ea(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs", "a:latin"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, shadow=False,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shadow:
        el = sp._element.spPr
        ef = el.makeelement(qn('a:effectLst'), {})
        sh = el.makeelement(qn('a:outerShdw'),
                {'blurRad':'80000','dist':'35000','dir':'5400000','rotWithShape':'0'})
        clr = el.makeelement(qn('a:srgbClr'), {'val':'1B2C38'})
        al  = el.makeelement(qn('a:alpha'), {'val':'20000'})
        clr.append(al); sh.append(clr); ef.append(sh); el.append(ef)
    return sp

def bar(slide, x, y, w, h, color):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp

def text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line_spacing=1.0, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    if isinstance(runs, tuple): paras = [[runs]]
    else:
        paras = [[p] if isinstance(p, tuple) else list(p) for p in runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        for seg in para:
            r = p.add_run(); r.text = seg[0]
            r.font.size = Pt(seg[1]); r.font.bold = seg[3]
            r.font.color.rgb = seg[2]; r.font.name = FONT
            _ea(r, FONT)
    return tb

def chip(slide, x, y, label, fill=TEAL, fg=WHITE, size=12.5):
    w = 0.18 + len(label) * (size/72.0) * 1.08
    sp = rect(slide, x, y, w, 0.42, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    tf = sp.text_frame; tf.word_wrap = False
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size=Pt(size); r.font.bold=True; r.font.color.rgb=fg; r.font.name=FONT
    _ea(r, FONT)
    return sp

PAGE = 0
def newslide(color=PAPER):
    global PAGE
    PAGE += 1
    s = prs.slides.add_slide(BLANK); bg(s, color); return s

def footer(slide, n):
    bar(slide, 0.9, 6.72, 11.53, 0.014, MINT2)
    text(slide, 11.35, 6.82, 1.08, 0.35, (f"{n:02d}", 10, GRAY, True),
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

def note(slide, s, x=0.9, y=6.82):
    text(slide, x, y, 9.8, 0.35, (s, 10, GRAY, False), anchor=MSO_ANCHOR.MIDDLE)

def head(slide, kick, title_runs, kfill=TEAL, sub=None):
    chip(slide, 0.9, 0.78, kick, fill=kfill)
    text(slide, 0.87, 1.34, 11.6, 1.05, title_runs, line_spacing=1.03)
    if sub:
        text(slide, 0.9, 2.32, 11.4, 0.55, (sub, 15.5, GRAY, False))

def say(slide, x, y, w, h, quote, fill=MINT, edge=TEAL, size=17, bold_parts=None):
    """원장이 그대로 말하는 문장 카드."""
    rect(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    bar(slide, x, y, 0.075, h, edge)
    runs = quote if isinstance(quote, list) else [[("“" + quote + "”", size, INK, False)]]
    text(slide, x+0.42, y, w-0.75, h, runs, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.32)

def section(num, title, sub):
    s = newslide(DEEP)
    bar(s, 0, 0, 13.333, 0.22, TEAL)
    bar(s, 0, 7.28, 13.333, 0.22, GOLD)
    text(s, 0.95, 1.9, 6, 2.2, (num, 140, RGBColor(0x18,0x5C,0x76), True))
    bar(s, 1.05, 4.05, 2.0, 0.055, GOLD)
    text(s, 1.0, 4.35, 11, 0.95, [[(title, 38, WHITE, True)]])
    text(s, 1.03, 5.35, 10.8, 0.7, (sub, 16.5, TEALL, False))
    return s


# =====================================================================
# 01 표지
# =====================================================================
s = newslide(DEEP)
bar(s, 0, 0, 13.333, 0.26, TEAL)
bar(s, 0, 7.24, 13.333, 0.26, GOLD)
c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.4), Inches(-1.7), Inches(5.8), Inches(5.8))
c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0x16,0x59,0x73); c.line.fill.background(); c.shadow.inherit=False
c2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.0), Inches(3.5), Inches(4.6), Inches(4.6))
c2.fill.solid(); c2.fill.fore_color.rgb = RGBColor(0x14,0x54,0x6D); c2.line.fill.background(); c2.shadow.inherit=False
chip(s, 0.95, 1.3, "의원 진료실 · 원장용", fill=TEAL)
text(s, 0.9, 2.1, 10.2, 2.4, [
    [("환자를 안심시키는", 50, WHITE, True)],
    [("시술 멘트 매뉴얼", 50, TEALL, True)],
], line_spacing=1.07)
bar(s, 0.98, 4.6, 2.1, 0.05, GOLD)
text(s, 0.95, 4.88, 10.2, 1.1, [
    [("말 한마디로 통증을 줄이고, 다시 오게 만듭니다", 19, WHITE, False)],
    [("척추 중재시술 · 관절 주사 · 곤란한 순간 응대", 14.5, TEALL, False)],
], line_spacing=1.4)

# =====================================================================
# 02 왜 이 매뉴얼인가
# =====================================================================
s = newslide(PAPER)
bar(s, 0, 0, 0.32, 7.5, TEAL)
chip(s, 1.2, 1.1, "이 매뉴얼을 쓰는 이유", fill=MINT2, fg=DEEP)
text(s, 1.15, 1.85, 11.2, 2.3, [
    [("같은 주사라도", 40, INK, True)],
    [("말에 따라 통증이 달라집니다", 40, TEAL, True)],
], line_spacing=1.08)
text(s, 1.2, 4.4, 10.8, 1.6, [
    [("시술 실력은 이미 갖추셨습니다. 남는 변수는 ", 18, INK, False),
     ("‘무슨 말을 하느냐’", 18, DEEP, True), ("입니다.", 18, INK, False)],
    [("환자가 기억하는 건 바늘이 아니라 ", 18, INK, False),
     ("그때 들은 말과 표정", 18, DEEP, True), ("입니다.", 18, INK, False)],
], line_spacing=1.5)
footer(s, PAGE)

# =====================================================================
# 03 세 가지 원칙
# =====================================================================
s = newslide(PAPER)
head(s, "전체 원칙", [[("이 세 가지만 지키면 됩니다", 32, INK, True)]])
cards = [
    ("1", "숨기지 않는다", "아픈 걸 아프다고 하되\n짧게, 담담하게", DEEP),
    ("2", "끝을 알려준다", "“몇 초면 지나가요”\n시간으로 닫는다", TEAL),
    ("3", "권한을 준다", "“불편하면 말씀하세요”\n멈출 수 있게 한다", GOLD),
]
for i,(n,t,d,col) in enumerate(cards):
    x = 0.9 + i*4.05
    rect(s, x, 2.85, 3.7, 3.25, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
    rect(s, x, 2.85, 3.7, 0.62, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    bar(s, x, 3.17, 3.7, 0.3, col)
    text(s, x, 2.85, 3.7, 0.62, (f"원칙 {n}", 13.5, WHITE, True), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x+0.38, 3.85, 3.0, 0.65, (t, 22, INK, True))
    text(s, x+0.4, 4.65, 3.0, 1.2, [[(l, 14.5, GRAY, False)] for l in d.split("\n")], line_spacing=1.3)
footer(s, PAGE)

# =====================================================================
# 04 섹션 1
# =====================================================================
section("01", "바늘 들어가는 순간", "가장 많이 실수하는 3초를 다듬습니다")

# =====================================================================
# 05 따끔은 그대로 쓰세요 (핵심 정정)
# =====================================================================
s = newslide(PAPER)
head(s, "먼저 정리", [[("“따끔”은 ", 33, INK, True), ("그대로 쓰셔도", 33, TEAL, True), (" 됩니다", 33, INK, True)]])
text(s, 0.9, 2.45, 11.3, 0.6, ("피부를 뚫는 건 실제로 따끔합니다. 이걸 ‘누르는 느낌’이라 하면 환자는 속았다고 느낍니다.", 16.5, INK, False))
rect(s, 0.9, 3.3, 11.5, 1.35, fill=CORALB, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
text(s, 1.35, 3.3, 3.3, 1.35, ("문제는 이것", 17, CORAL, True), anchor=MSO_ANCHOR.MIDDLE)
text(s, 3.9, 3.3, 8.2, 1.35, [
    [("‘따끔’이라는 단어가 아니라 ", 16, INK, False), ("아픔을 부풀리고 물고 늘어지는 것", 16, CORAL, True)],
    [("“많이 아파요” · “여기가 제일 아파요” · “벌에 쏘이듯”", 15, GRAY, False)],
], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
rect(s, 0.9, 4.85, 11.5, 1.35, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
text(s, 1.35, 4.85, 3.3, 1.35, ("그래서 이렇게", 17, TEAL, True), anchor=MSO_ANCHOR.MIDDLE)
text(s, 3.9, 4.85, 8.2, 1.35, [
    [("한 번만 담담하게 + ", 16, INK, False), ("시간으로 닫고", 16, DEEP, True), (" + ", 16, INK, False),
     ("멈출 권한", 16, DEEP, True), ("을 준다", 16, INK, False)],
    [("“잠깐 따끔하고 몇 초면 지나가요. 불편하면 말씀하세요.”", 15, GRAY, False)],
], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
footer(s, PAGE)

# =====================================================================
# 06 감각을 단계로 나눠 말한다
# =====================================================================
s = newslide(PAPER)
head(s, "정확하게 말하기", [[("감각은 ", 32, INK, True), ("단계마다 다릅니다", 32, TEAL, True)]],
     sub="한 단어로 뭉치지 말고, 그 순간의 실제 감각을 그대로 알려주세요.")
rows = [
    ("피부 뚫을 때", "따끔", "“잠깐 따끔하고 몇 초면 지나가요”", TEAL),
    ("마취약 들어갈 때", "시큰·화끈", "“약이 퍼지면서 시큰한데 금방 무뎌집니다”", DEEP),
    ("바늘 진행할 때", "누르는 느낌", "“이제부턴 지그시 누르는 느낌이에요”", TEAL),
    ("약 들어갈 때", "뻐근·차오름", "“묵직하게 차오르는 느낌은 정상이에요”", DEEP),
]
for i,(when, feel, line, col) in enumerate(rows):
    y = 3.05 + i*0.87
    rect(s, 0.9, y, 11.5, 0.76, fill=WHITE if i%2==0 else MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1, shadow=(i%2==0))
    text(s, 1.25, y, 2.4, 0.76, (when, 14.5, GRAY, True), anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.75, y, 1.9, 0.76, (feel, 16, col, True), anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.8, y, 6.4, 0.76, (line, 15.5, INK, False), anchor=MSO_ANCHOR.MIDDLE)
note(s, "‘찌릿’은 아껴두세요 — 신경근차단에서 “다리로 찌릿하면 말씀하세요”라는 안전 신호로 써야 합니다.")
footer(s, PAGE)

# =====================================================================
# 07 3초 스크립트
# =====================================================================
s = newslide(PAPER)
head(s, "바로 쓰는 3초", [[("찌르기 직전, 이 한 문장", 32, INK, True)]])
rect(s, 0.9, 2.6, 11.5, 1.55, fill=DEEP, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, shadow=True)
text(s, 1.4, 2.6, 10.6, 1.55, [
    [("“자, 시작할게요. ", 21, WHITE, False), ("잠깐 따끔", 21, GOLD, True),
     ("하고 ", 21, WHITE, False), ("몇 초면 지나가요", 21, GOLD, True), (".", 21, WHITE, False)],
    [("  셋 셀게요 — 하나, 둘, 셋.”", 21, WHITE, False)],
], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.35)
why = [("‘자, 시작할게요’", "예고 — 놀라지 않게"),
       ("‘잠깐 따끔’", "정직 — 속이지 않게"),
       ("‘몇 초면 지나가요’", "끝 — 견딜 수 있게"),
       ("‘하나, 둘, 셋’", "통제 — 대비할 수 있게")]
for i,(a,b) in enumerate(why):
    x = 0.9 + (i%2)*5.85
    y = 4.5 + (i//2)*0.88
    rect(s, x, y, 5.5, 0.76, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, x+0.35, y, 2.5, 0.76, (a, 14.5, DEEP, True), anchor=MSO_ANCHOR.MIDDLE)
    text(s, x+2.95, y, 2.3, 0.76, (b, 13.5, GRAY, False), anchor=MSO_ANCHOR.MIDDLE)
footer(s, PAGE)

# =====================================================================
# 08 시술 중
# =====================================================================
s = newslide(PAPER)
head(s, "시술 중", [[("침묵을 두지 마세요", 32, INK, True)]],
     sub="아무 말 없이 3초가 지나면 환자는 뭔가 잘못됐다고 생각합니다.")
qs = [
    "지금 잘 들어가고 있어요. 절반쯤 지났습니다.",
    "누르는 느낌은 정상이에요. 날카롭게 오면 바로 말씀해 주세요.",
    "숨 편하게 쉬세요. 참지 마시고요.",
    "거의 다 됐어요. 조금만요.",
]
for i,q in enumerate(qs):
    y = 3.0 + i*0.92
    say(s, 0.9, y, 11.5, 0.78, q, fill=MINT if i%2==0 else WHITE, edge=TEAL, size=16.5)
    if i%2==1:
        rect(s, 0.9, y, 11.5, 0.78, fill=None, line=MINT2, line_w=1.2,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
note(s, "진행 상황을 말해주는 것 자체가 진통 효과가 있습니다. 특히 ‘절반 지났다’는 끝이 보이게 합니다.")
footer(s, PAGE)

# =====================================================================
# 09 시술 후 30초
# =====================================================================
s = newslide(PAPER)
head(s, "시술 후 30초", [[("여기서 ", 32, INK, True), ("재방문이 갈립니다", 32, TEAL, True)]],
     sub="시술이 끝난 직후 30초가 그날 진료 전체의 인상을 결정합니다.")
steps = [
    ("① 인정", "“잘 참으셨어요. 생각보다 금방 끝났죠?”"),
    ("② 확인", "“지금 어지럽거나 불편한 데는 없으세요?”"),
    ("③ 예고", "“오늘 저녁이나 내일 잠깐 뻐근할 수 있는데 정상이에요.”"),
    ("④ 연결", "“이상하면 참지 마시고 바로 전화 주세요.”"),
]
for i,(t,q) in enumerate(steps):
    y = 3.05 + i*0.9
    rect(s, 0.9, y, 2.0, 0.76, fill=DEEP if i%2==0 else TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 0.9, y, 2.0, 0.76, (t, 15, WHITE, True), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 3.05, y, 9.35, 0.76, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 3.45, y, 8.7, 0.76, (q, 16, INK, False), anchor=MSO_ANCHOR.MIDDLE)
note(s, "④번을 빠뜨리면 환자는 불안할 때 우리 대신 인터넷을 찾습니다.")
footer(s, PAGE)

# =====================================================================
# 10 하지 말아야 할 말
# =====================================================================
s = newslide(PAPER)
head(s, "금지어", [[("이 말만은 ", 32, INK, True), ("빼주세요", 32, CORAL, True)]], kfill=CORAL)
bads = [
    ("“여기가 제일 아픈 부분이에요”", "아픔에 주의를 못 박습니다"),
    ("“하나도 안 아파요”", "거짓말은 다음 시술을 어렵게 합니다"),
    ("“많이 아팠죠?” (끝난 뒤)", "지나간 통증을 다시 각인시킵니다"),
    ("“절대 움직이면 안 돼요”", "긴장으로 오히려 더 아파집니다"),
    ("“어? 이상하네” · 혼잣말", "환자는 최악을 상상합니다"),
]
for i,(b,why) in enumerate(bads):
    y = 2.62 + i*0.82
    rect(s, 0.9, y, 11.5, 0.7, fill=CORALB if i%2==0 else WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 1.3, y, 0.5, 0.7, ("×", 18, CORAL, True), anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.85, y, 5.6, 0.7, (b, 16, INK, False), anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.6, y, 4.5, 0.7, (why, 13.5, GRAY, False), anchor=MSO_ANCHOR.MIDDLE)
footer(s, PAGE)

# =====================================================================
# 11 섹션 2
# =====================================================================
section("02", "환자가 꼭 묻는 질문", "말문 막히지 않는 표준 답변")

# ---- 질문 응대 공통 레이아웃 ----
def qna(kick, question, answers, avoid=None, tip=None):
    s = newslide(PAPER)
    chip(s, 0.9, 0.78, kick, fill=GOLD)
    text(s, 0.87, 1.32, 11.6, 1.0, [[("“" + question + "”", 30, INK, True)]], line_spacing=1.05)
    y = 2.62
    for a in answers:
        h = 1.0 if len(a) > 46 else 0.8
        say(s, 0.9, y, 11.5, h, a, fill=MINT, edge=TEAL, size=16.5)
        y += h + 0.22
    if avoid:
        rect(s, 0.9, y, 11.5, 0.72, fill=CORALB, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        text(s, 1.3, y, 0.5, 0.72, ("×", 17, CORAL, True), anchor=MSO_ANCHOR.MIDDLE)
        text(s, 1.85, y, 10.2, 0.72, (avoid, 15.5, INK, False), anchor=MSO_ANCHOR.MIDDLE)
        y += 0.9
    if tip: note(s, tip)
    footer(s, PAGE)
    return s

# 12
qna("환자 질문 · 오해", "이거 뼈주사죠?",
    ["뼈에 놓는 주사가 아니라, 신경 옆 정확한 지점에 약을 넣는 겁니다. 엑스레이로 위치를 보면서 해요.",
     "말씀하시는 ‘뼈주사’는 보통 스테로이드를 말씀하시는 건데, 오늘은 양을 최소로 쓰고 자주 반복하지 않습니다."],
    avoid="“뼈주사 아니에요” 하고 끝내기 — 환자는 여전히 못 미더워합니다.",
    tip="부정만 하지 말고 ‘그럼 뭔데요?’에 대한 답까지 한 문장으로 주세요.")

# 13
qna("환자 질문 · 오해", "스테로이드 많이 맞으면 안 좋다던데요",
    ["맞습니다. 그래서 저희도 횟수를 정해두고 씁니다. 무제한으로 놓는 약이 아니에요.",
     "오늘 쓰는 양은 먹는 약으로 치면 아주 적은 양이고, 간격을 충분히 둘 겁니다."],
    avoid="“그런 건 인터넷 얘기예요” — 환자의 걱정을 무시하면 신뢰가 깎입니다.",
    tip="걱정을 먼저 인정하면(“맞습니다”) 그 다음 설명이 훨씬 잘 들어갑니다.")

# 14
qna("환자 질문 · 기대치", "이거 맞으면 낫는 거죠?",
    ["이 주사는 통증을 줄여서 움직일 수 있게 해주는 역할이에요. 그 사이에 운동이랑 자세를 잡는 게 진짜 치료입니다.",
     "효과는 사람마다 달라서, 오늘 맞아보고 다음에 오셨을 때 얼마나 편했는지 알려주시면 계획을 정하겠습니다."],
    avoid="“한 방이면 끝나요” — 안 나으면 그대로 불신으로 돌아옵니다.",
    tip="확답 대신 ‘함께 확인하자’는 프레임이 재방문을 만듭니다.")

# 15
qna("환자 질문 · 기대치", "몇 번이나 맞아야 돼요?",
    ["보통 한 번 맞고 반응을 봅니다. 좋으면 굳이 더 안 맞아도 되고요.",
     "계속 맞는 치료가 아니라, 필요할 때만 쓰는 겁니다. 오늘 맞고 2주쯤 뒤에 어떤지 보고 정하시죠."],
    avoid="“3번은 맞으셔야 돼요” — 처음부터 횟수를 못 박으면 상술로 보입니다.",
    tip="‘반응 보고 정한다’가 가장 정직하면서 신뢰도 높은 답입니다.")

# 16
qna("환자 질문 · 안전", "방사선 많이 쬐는 거 아니에요?",
    ["위치 확인할 때만 잠깐씩 쓰고, 다 합쳐도 몇십 초 정도예요.",
     "필요한 만큼만 최소로 쓰고, 정확히 놓기 위해서 쓰는 거라 오히려 더 안전합니다."],
    avoid="“그 정도는 괜찮아요” — 숫자 없이 넘어가면 안심이 안 됩니다.")

# 17
qna("환자 질문 · 안전", "신경 건드리면 마비되는 거 아니에요?",
    ["바늘이 향하는 건 신경이 아니라 뼈 위의 정해진 지점이에요. 뼈에 닿으면 거기서 멈춥니다.",
     "그리고 깨어 계시니까, 다리로 찌릿한 느낌이 오면 바로 말씀해 주시면 제가 위치를 조정합니다."],
    avoid="위험을 길게 나열하기 — 동의서에 있는 내용을 다시 읊으면 불안만 커집니다.",
    tip="‘환자가 말해주는 것이 안전장치’라고 하면 불안이 참여로 바뀝니다.")

# =====================================================================
# 18 섹션 3
# =====================================================================
section("03", "시술별 한마디", "그 시술에서 꼭 해야 하는 말만")

# ---- 시술별 카드 2개 배치 ----
def proc2(kick, items):
    s = newslide(PAPER)
    chip(s, 0.9, 0.78, kick, fill=DEEP)
    text(s, 0.87, 1.32, 11.6, 0.95, [[(items[0][0] + "  ·  " + items[1][0], 30, INK, True)]])
    for i,(name, must, lines) in enumerate(items):
        x = 0.9 + i*5.85
        rect(s, x, 2.5, 5.5, 3.75, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, shadow=True)
        rect(s, x, 2.5, 5.5, 0.68, fill=DEEP if i==0 else TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        bar(s, x, 2.85, 5.5, 0.33, DEEP if i==0 else TEAL)
        text(s, x, 2.5, 5.5, 0.68, (name, 17, WHITE, True), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x+0.34, 3.35, 4.9, 0.45, (must, 13, GOLD, True))
        yy = 3.85
        for l in lines:
            text(s, x+0.34, yy, 4.92, 1.0, [[("“" + l + "”", 13.5, INK, False)]], line_spacing=1.28)
            yy += 0.95
    footer(s, PAGE)
    return s

# 19
proc2("척추", [
    ("경추 내측지신경차단", "꼭 할 말 — 반복 촬영은 안전 절차",
     ["편하게 누워서 하니 얘기하시면서 하셔도 돼요.",
      "여러 번 찍는 건 정확히 보려고 하는 겁니다.",
      "어지럽거나 식은땀 나면 바로 말씀하세요."]),
    ("흉추·요추 내측지신경차단", "꼭 할 말 — 표적은 신경이 아니라 뼈",
     ["엎드려서 하고, 배에 베개 받쳐드릴게요.",
      "다리로 뻗치는 시술이 아니라 누르는 느낌 위주예요.",
      "바늘은 뼈에 닿으면 거기서 멈춥니다."]),
])

# 20
proc2("척추", [
    ("신경근차단", "꼭 할 말 — 찌릿하면 ‘말로만’",
     ["다리로 찌릿한 느낌이 잠깐 올 수 있어요.",
      "위치가 잘 왔다는 신호니 걱정 안 하셔도 됩니다.",
      "그때 움직이지 마시고 말로만 알려주세요."]),
    ("경막외 차단", "꼭 할 말 — 응급 신호는 미리",
     ["등을 고양이처럼 말아주시면 수월해요.",
      "밀리는 느낌은 정상, 찌릿하면 말씀하세요.",
      "힘 빠짐이 심해지거나 소변이 안 나오면 연락 주세요."]),
])

# 21
proc2("관절", [
    ("무릎 관절강내", "꼭 할 말 — 몇 주짜리 다리라는 것",
     ["다리 힘 빼세요. 긴장하면 더 뻐근합니다.",
      "약 들어가며 차오르는 느낌은 정상이에요.",
      "이 주사가 버는 시간에 운동·체중을 같이 잡으셔야 해요."]),
    ("어깨 (견봉하)", "꼭 할 말 — 효과는 2~3일 뒤부터",
     ["뒤에서 놓아 바늘은 안 보이실 거예요. 힘 빼세요.",
      "어지럽거나 메스꺼우면 바로 말씀하세요.",
      "약효는 2~3일 뒤부터라 바로 안 좋아져도 정상이에요."]),
])

# 22
s = newslide(PAPER)
chip(s, 0.9, 0.78, "관절", fill=DEEP)
text(s, 0.87, 1.32, 11.6, 0.95, [[("팔꿈치 (외측상과염)", 30, INK, True)]])
text(s, 0.9, 2.35, 11.4, 0.6, ("단기 효과는 좋지만 장기적으로는 재발이 많은 시술입니다. 이 사실을 미리 말해두면 나중에 신뢰를 지킵니다.", 15.5, GRAY, False))
qs = [
    "이 부위는 원래 예민해서 좀 뻐근하실 거예요. 최대한 천천히 놓겠습니다.",
    "이 주사는 지금 아픈 걸 줄여주는 거고, 근본은 스트레칭이랑 근력운동이에요.",
    "주사만 반복하면 오히려 나중에 더 안 좋을 수 있어서, 오늘 맞고 운동을 시작하시는 게 중요합니다.",
    "오늘 밤이나 내일 잠깐 더 아플 수 있는데 2~3일이면 가라앉습니다.",
]
for i,q in enumerate(qs):
    y = 3.15 + i*0.92
    say(s, 0.9, y, 11.5, 0.78, q, fill=MINT if i%2==0 else WHITE, edge=TEAL, size=16)
    if i%2==1:
        rect(s, 0.9, y, 11.5, 0.78, fill=None, line=MINT2, line_w=1.2,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
footer(s, PAGE)

# =====================================================================
# 23 섹션 4
# =====================================================================
section("04", "곤란한 순간", "당황하면 신뢰가 무너지는 다섯 장면")

# ---- 상황 대응 레이아웃 ----
def scene(kick, title, situation, dos, tip=None):
    s = newslide(PAPER)
    chip(s, 0.9, 0.78, kick, fill=CORAL)
    text(s, 0.87, 1.32, 11.6, 0.95, [[(title, 30, INK, True)]])
    rect(s, 0.9, 2.35, 11.5, 0.72, fill=CORALB, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 1.3, 2.35, 10.8, 0.72, (situation, 15.5, INK, False), anchor=MSO_ANCHOR.MIDDLE)
    y = 3.35
    for d in dos:
        h = 0.95 if len(d) > 46 else 0.78
        say(s, 0.9, y, 11.5, h, d, fill=MINT, edge=TEAL, size=16.5)
        y += h + 0.2
    if tip: note(s, tip)
    footer(s, PAGE)
    return s

# 24
scene("곤란한 순간 ①", "환자가 어지럽다고 합니다",
      "주사 후 실신은 대부분 15분 안에 생깁니다. 앉힌 채로 버티게 하면 넘어집니다.",
      ["어지러우시군요. 괜찮습니다, 바로 눕혀드릴게요.",
       "다리를 이렇게 올릴게요. 금방 편해지실 거예요. 천천히 숨 쉬세요.",
       "이제 혈색 돌아왔어요. 바로 일어나지 마시고 몇 분 더 누워 계세요."],
      tip="당황한 티를 내지 않는 것이 핵심입니다. ‘가끔 있는 일’이라고 말해주면 환자도 진정합니다.")

# 25
scene("곤란한 순간 ②", "바늘을 무서워합니다",
      "바늘 공포는 아주 흔합니다. 다그치면 그날 시술도, 다음 방문도 사라집니다.",
      ["바늘 무서워하시는 분 정말 많아요. 부끄러워하실 것 하나도 없습니다.",
       "안 보셔도 돼요. 저쪽 보시면서 저랑 얘기해요.",
       "준비되면 말씀하세요. 준비되실 때까지 기다릴게요."],
      tip="‘기다려주겠다’는 말이 통제권을 돌려줍니다. 실제로는 몇 초 안 걸립니다.")

# 26
scene("곤란한 순간 ③", "다음 날 더 아프다고 전화가 옵니다",
      "주사 후 일시적 통증 악화는 흔한 반응입니다. 미리 말해뒀다면 컴플레인이 아니라 확인 전화가 됩니다.",
      ["네, 그럴 수 있다고 말씀드렸던 반응이에요. 보통 2~3일이면 가라앉습니다.",
       "얼음찜질 하루 두세 번 해보시고, 무리한 사용만 피해주세요.",
       "다만 붓고 벌겋게 열이 나거나 통증이 계속 심해지면 바로 오세요."],
      tip="시술 직후에 미리 예고해두면 이 전화 자체가 크게 줄어듭니다. 09번 슬라이드 ③번을 꼭 하세요.")

# 27
scene("곤란한 순간 ④", "효과가 없다고 다시 왔습니다",
      "여기서 방어적으로 나가면 환자를 잃습니다. 정보로 받아들이면 다음 치료가 정확해집니다.",
      ["그러셨군요. 언제부터 다시 아프기 시작했는지 알려주시겠어요?",
       "잠깐이라도 편했던 시간이 있었다면 그것도 중요한 정보예요.",
       "그럼 위치를 다시 보고, 다른 방법도 같이 생각해보겠습니다."],
      tip="‘효과 없음’도 진단 정보입니다. 이렇게 다루면 실패가 신뢰로 바뀝니다.")

# 28
scene("곤란한 순간 ⑤", "보호자가 옆에서 불안해합니다",
      "보호자의 불안은 환자에게 그대로 옮겨갑니다. 보호자를 먼저 안심시켜야 합니다.",
      ["보호자분, 옆에서 손 잡아주시면 큰 도움이 됩니다.",
       "지금 하는 건 위치를 확인하는 과정이라 시간이 좀 걸립니다. 정상이에요.",
       "끝나면 어떻게 지내셔야 하는지 같이 설명드릴게요."],
      tip="보호자에게 역할을 주면 불안이 협조로 바뀝니다.")

# =====================================================================
# 29 한 장 요약
# =====================================================================
s = newslide(PAPER)
head(s, "한 장 요약", [[("오늘부터 이것만", 32, INK, True)]])
summary = [
    ("찌르기 전", "“잠깐 따끔하고 몇 초면 지나가요. 셋 셀게요.”", TEAL),
    ("바늘 진행", "“누르는 느낌은 정상이에요. 날카로우면 말씀하세요.”", DEEP),
    ("시술 중", "“절반쯤 지났어요. 잘하고 계세요.”", TEAL),
    ("끝난 직후", "“잘 참으셨어요. 어지럽거나 불편한 데 없으세요?”", DEEP),
    ("보내기 전", "“내일 잠깐 뻐근할 수 있어요. 이상하면 바로 전화 주세요.”", GOLD),
]
for i,(when, line, col) in enumerate(summary):
    y = 2.5 + i*0.85
    rect(s, 0.9, y, 2.15, 0.72, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    text(s, 0.9, y, 2.15, 0.72, (when, 14.5, WHITE, True), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 3.2, y, 9.2, 0.72, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    text(s, 3.6, y, 8.6, 0.72, (line, 16, INK, False), anchor=MSO_ANCHOR.MIDDLE)
footer(s, PAGE)

# =====================================================================
# 30 마무리
# =====================================================================
s = newslide(DEEP)
bar(s, 0, 0, 13.333, 0.26, TEAL)
bar(s, 0, 7.24, 13.333, 0.26, GOLD)
chip(s, 0.95, 1.05, "기억할 것", fill=TEAL)
text(s, 0.9, 1.8, 11.5, 2.0, [
    [("아픈 걸 숨기지 말고,", 42, WHITE, True)],
    [("끝을 알려주세요.", 42, TEALL, True)],
], line_spacing=1.08)
bar(s, 0.98, 4.15, 2.1, 0.05, GOLD)
text(s, 0.95, 4.55, 11.2, 1.5, [
    [("환자는 아프지 않아서가 아니라, ", 19, WHITE, False),
     ("아플 때 곁에 있어줘서", 19, GOLD, True), (" 다시 옵니다.", 19, WHITE, False)],
], line_spacing=1.4)
text(s, 0.95, 5.75, 11.2, 0.9, [
    [("근거: 시술 직전 위협적 표현은 통증을 높이고(Varelmann 2010), 경고성 멘트는 통증·불안을 함께 올립니다(Lang 2005).", 12, TEALL, False)],
    [("반면 ‘따끔’ 같은 중립적 예고 자체는 통증을 높이지 않았습니다(Vijayan 2015).", 12, TEALL, False)],
], line_spacing=1.35)

prs.save("/home/user/ppt-work/시술멘트_매뉴얼_의사용.pptx")
print("OK — slides:", len(prs.slides._sldIdLst), "PAGE:", PAGE)
